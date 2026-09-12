from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from q3.config import (
    FOUR_HOUR_BLOCKS,
    REQUIRED_TRAJECTORY_COLS,
    RESULT3_EXPORT_N_DAYS,
    RESULT3_TEMPLATE,
    T,
    YEAR_END_SOC_A_KWH,
)
from q3.export_result3 import (
    compare_to_annual_summary,
    enrich_dispatch,
    export_dates,
    inspect_result3_workbook,
    locked_or_mutable_label,
    write_result3_xlsx,
)
from q3.pilot import DayRun


def _trajectory(date: str, seed: int, emergency_at: int | None = None) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    g0 = rng.uniform(20.0, 80.0, size=T)
    gf = g0.copy()
    gf[72:] = g0[72:] * 0.9
    last_update = np.array(["00:00"] * 36 + ["06:00"] * 36 + ["12:00"] * 36 + ["18:00"] * 36)
    soc_start = 6000.0 + seed
    charge = rng.uniform(0.0, 5.0, size=T)
    discharge = rng.uniform(0.0, 4.0, size=T)
    soc_end = soc_start + np.cumsum(charge - discharge)
    emergency = np.zeros(T)
    if emergency_at is not None:
        emergency[emergency_at] = 3.5
    price = np.full(T, 0.4)
    phi = price * gf + 0.5 * price * np.abs(gf - g0)
    starts = np.empty(T)
    starts[0] = soc_start
    starts[1:] = soc_end[:-1]
    return pd.DataFrame(
        {
            "date": date,
            "period_index": np.arange(T),
            "time_label": [f"t{t}" for t in range(T)],
            "g0_kwh": g0,
            "g_final_kwh": gf,
            "x_kwh": gf,
            "emergency_kwh": emergency,
            "charge_kwh": charge,
            "discharge_kwh": discharge,
            "curtailment_kwh": rng.uniform(0.0, 1.0, size=T),
            "soc_start_kwh": starts,
            "soc_end_kwh": soc_end,
            "actual_load_kwh": rng.uniform(30.0, 90.0, size=T),
            "actual_pv_kwh": rng.uniform(0.0, 20.0, size=T),
            "locked_or_mutable": [
                locked_or_mutable_label(t, last_update[t]) for t in range(T)
            ],
            "last_update_time": last_update,
            "price": price,
            "phi_yuan": phi,
            "emergency_cost_yuan": 5.0 * price * emergency,
            "run_id": f"{date}__M1_M6",
            "year_end_boundary": "A_q2_aligned",
            "year_end_soc_kwh": YEAR_END_SOC_A_KWH,
            "strategy": "M1_M6",
            "balance_residual_kwh": np.zeros(T),
        }
    )


class LockedLabelTests(unittest.TestCase):
    def test_update_prefix_is_locked(self) -> None:
        self.assertEqual(locked_or_mutable_label(10, "00:00"), "mutable")
        self.assertEqual(locked_or_mutable_label(35, "06:00"), "locked")
        self.assertEqual(locked_or_mutable_label(36, "06:00"), "mutable")
        self.assertEqual(locked_or_mutable_label(107, "18:00"), "locked")
        self.assertEqual(locked_or_mutable_label(108, "18:00"), "mutable")


class EnrichDispatchTests(unittest.TestCase):
    def test_soc_start_and_required_columns(self) -> None:
        dispatch = pd.DataFrame(
            {
                "time": [f"{t:02d}:10" for t in range(T)],
                "period": np.arange(T),
                "planned_g0_kwh": np.ones(T),
                "final_g_kwh": np.ones(T),
                "normal_x_kwh": np.ones(T),
                "load_kwh": np.full(T, 2.0),
                "actual_pv_kwh": np.zeros(T),
                "charge_kwh": np.zeros(T),
                "discharge_kwh": np.zeros(T),
                "curtailment_kwh": np.zeros(T),
                "emergency_kwh": np.zeros(T),
                "soc_kwh": 6000.0 + np.arange(T),
                "last_update_time": ["00:00"] * T,
                "price": np.full(T, 0.5),
                "phi_yuan": np.full(T, 0.5),
                "emergency_cost_yuan": np.zeros(T),
                "balance_residual_kwh": np.zeros(T),
            }
        )
        run = DayRun(
            date="2025-02-01",
            strategy="M1_M6",
            load_information_case="causal_load_main",
            pv_mapping_mode="linear_anchor_main",
            settlement_mode="anchor_final_main",
            with_terminal_value=True,
            run_id="x",
            dispatch=dispatch,
            update_log=pd.DataFrame(),
            ledger=pd.DataFrame(),
            summary={"soc_start_kwh": 6000.0, "year_end_boundary": "A_q2_aligned"},
        )
        frame = enrich_dispatch(run)
        for col in REQUIRED_TRAJECTORY_COLS:
            self.assertIn(col, frame.columns)
        self.assertAlmostEqual(float(frame["soc_start_kwh"].iloc[0]), 6000.0)
        self.assertAlmostEqual(float(frame["soc_start_kwh"].iloc[1]), float(frame["soc_end_kwh"].iloc[0]))
        self.assertEqual(len(frame), T)


@unittest.skipUnless(RESULT3_TEMPLATE.exists(), "result3 template unavailable")
class Result3WorkbookTests(unittest.TestCase):
    def test_template_fill_is_unrotated_and_excludes_january(self) -> None:
        dates = export_dates()
        self.assertEqual(len(dates), RESULT3_EXPORT_N_DAYS)
        self.assertEqual(dates[0], "2025-02-01")
        self.assertEqual(dates[-1], "2025-12-31")
        by_date = {}
        for i, date in enumerate(dates):
            emergency_at = 5 if date == "2025-02-01" else None
            frame = _trajectory(date, seed=i + 1, emergency_at=emergency_at)
            # Distinct marker in period 0 so rotation would be visible.
            frame.loc[0, "g0_kwh"] = 111.0 + i
            frame.loc[0, "g_final_kwh"] = 111.0 + i
            by_date[date] = frame
        with tempfile.TemporaryDirectory() as tmp:
            dest = Path(tmp) / "result3.xlsx"
            meta = write_result3_xlsx(by_date, dest=dest, source_commit="test")
            inspect = inspect_result3_workbook(dest)
            self.assertEqual(inspect["n_anomalies"], 0, inspect["anomalies"])
            self.assertEqual(inspect["n_plan_dates"], 334)
            self.assertEqual(inspect["n_emergency_rows"], 1)
            self.assertTrue(inspect["has_notes_sheet"])
            wb = load_workbook(dest)
            plan = wb["计划购电量"]
            adjust = wb["调整购电量"]
            charge = wb["充放电量"]
            notes = wb["说明"]
            self.assertEqual(plan.cell(2, 1).value, datetime(2025, 2, 1))
            self.assertEqual(plan.cell(plan.max_row, 1).value, datetime(2025, 12, 31))
            self.assertAlmostEqual(float(plan.cell(2, 2).value), 111.0)
            self.assertAlmostEqual(float(adjust.cell(2, 2).value), 111.0)
            # Unadjusted prefix equals g0 and is not blank.
            for col in range(2, 2 + 72):
                self.assertIsNotNone(adjust.cell(2, col).value)
                self.assertAlmostEqual(
                    float(adjust.cell(2, col).value), float(plan.cell(2, col).value)
                )
            self.assertEqual(plan.cell(1, 2).value, "0:10-0:20")
            self.assertEqual(plan.cell(1, 145).value, "0:00-0:10+1")
            self.assertNotIn("January", "".join(str(c.value) for c in notes["A"] if c.value))
            self.assertIn("M1_M6", notes.cell(2, 1).value)
            self.assertIn("1200", notes.cell(3, 1).value)
            # Charge sheet: date on first row of day, SOC on first two rows only.
            self.assertEqual(charge.cell(2, 1).value, datetime(2025, 2, 1))
            self.assertIsNone(charge.cell(3, 1).value)
            self.assertEqual(charge.cell(2, 2).value, FOUR_HOUR_BLOCKS[0][0])
            first = by_date["2025-02-01"]
            self.assertAlmostEqual(
                float(charge.cell(2, 3).value), float(first["charge_kwh"].iloc[0:24].sum())
            )
            self.assertAlmostEqual(
                float(charge.cell(2, 6).value), float(first["soc_start_kwh"].iloc[0])
            )
            self.assertAlmostEqual(
                float(charge.cell(3, 6).value), float(first["soc_end_kwh"].iloc[-1])
            )
            self.assertIsNone(charge.cell(4, 6).value)
            self.assertEqual(charge.max_row, 1 + 334 * 6)
            em = wb["紧急购电量"]
            self.assertEqual(em.cell(2, 1).value, datetime(2025, 2, 1))
            self.assertAlmostEqual(float(em.cell(2, 3).value), 3.5)
            self.assertEqual(meta["n_export_days"], 334)
            self.assertNotIn("#REF!", str(plan.cell(2, 2).value))


class SummaryCompareTests(unittest.TestCase):
    def test_solver_tolerance_match(self) -> None:
        dates = ["2025-01-01", "2025-02-01"]
        rerun = pd.DataFrame(
            [
                {
                    "date": dates[0],
                    "total_cost_yuan": 10.0,
                    "emergency_kwh": 1.0,
                    "curtailment_kwh": 2.0,
                    "soc_start_kwh": 6000.0,
                    "soc_end_kwh": 6100.0,
                },
                {
                    "date": dates[1],
                    "total_cost_yuan": 20.0,
                    "emergency_kwh": 0.0,
                    "curtailment_kwh": 3.0,
                    "soc_start_kwh": 6100.0,
                    "soc_end_kwh": 1200.0,
                },
            ]
        )
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "q3_annual_daily_summary.csv"
            archive = rerun.copy()
            archive["strategy"] = "M1_M6"
            archive["year_end_boundary"] = "A_q2_aligned"
            archive["total_cost_yuan"] = archive["total_cost_yuan"] + 1e-6
            archive.to_csv(path, index=False)
            result = compare_to_annual_summary(rerun, path)
            self.assertTrue(result["all_match"])


if __name__ == "__main__":
    unittest.main()
