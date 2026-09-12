"""Export the official Q3 workbook from a single A + M1_M6 10-minute trajectory.

January is solved only for SOC continuity. The contest workbook covers
2025-02-01 through 2025-12-31. Does not reconstruct intervals from daily
summaries and does not rerun the other annual paths.
"""

from __future__ import annotations

import json
import shutil
import subprocess
from datetime import datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from q3.config import (
    ANNUAL_END,
    ANNUAL_OUTPUT_DIR,
    ANNUAL_START,
    FOUR_HOUR_BLOCKS,
    HOUR_TO_FIRST_MUTABLE,
    REQUIRED_TRAJECTORY_COLS,
    RESULT3_ANNUAL_COST_YUAN,
    RESULT3_EXPORT_AUDIT,
    RESULT3_EXPORT_END,
    RESULT3_EXPORT_N_DAYS,
    RESULT3_EXPORT_START,
    RESULT3_OFFICIAL_STRATEGY,
    RESULT3_TEMPLATE,
    RESULT3_XLSX,
    ROOT,
    T,
    YEAR_END_BOUNDARY_A,
    YEAR_END_SOC_A_KWH,
    dispatch_stem,
)
from q3.pilot import DayRun

PURCHASE_SHEET = "计划购电量"
ADJUST_SHEET = "调整购电量"
CHARGE_SHEET = "充放电量"
EMERGENCY_SHEET = "紧急购电量"
NOTES_SHEET = "说明"
N_TIME_COLS = T
COL_FIRST_TIME = 2
COL_DAY_QTY = 146
COL_DAY_FEE = 147
COST_MATCH_TOL_YUAN = 5e-3
ENERGY_MATCH_TOL_KWH = 1e-3
SOC_MATCH_TOL_KWH = 1e-3
EMERGENCY_EPS_KWH = 1e-8

TRAJECTORY_WRITE_COLS = list(REQUIRED_TRAJECTORY_COLS) + [
    "price",
    "phi_yuan",
    "emergency_cost_yuan",
    "run_id",
    "year_end_boundary",
    "year_end_soc_kwh",
    "strategy",
    "balance_residual_kwh",
]


def git_hash() -> str:
    try:
        return subprocess.check_output(
            ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
        ).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def export_dates() -> list[str]:
    return [
        ts.strftime("%Y-%m-%d")
        for ts in pd.date_range(RESULT3_EXPORT_START, RESULT3_EXPORT_END, freq="D")
    ]


def annual_calendar() -> list[str]:
    return [
        ts.strftime("%Y-%m-%d")
        for ts in pd.date_range(ANNUAL_START, ANNUAL_END, freq="D")
    ]


def main_dispatch_stem(date: str) -> str:
    return dispatch_stem(
        date,
        RESULT3_OFFICIAL_STRATEGY,
        "causal_load_main",
        "linear_anchor_main",
        "anchor_final_main",
        True,
        YEAR_END_SOC_A_KWH,
    )


def dispatch_path(date: str, dispatch_dir: Path | None = None) -> Path:
    folder = dispatch_dir or (ANNUAL_OUTPUT_DIR / "dispatch_daily")
    return folder / f"{main_dispatch_stem(date)}.csv"


def _update_hour(last_update_time: object) -> int:
    text = str(last_update_time).strip()
    hour = int(text.split(":")[0])
    if hour not in HOUR_TO_FIRST_MUTABLE:
        raise ValueError(f"unknown last_update_time {last_update_time!r}")
    return hour


def locked_or_mutable_label(period_index: int, last_update_time: object) -> str:
    first = HOUR_TO_FIRST_MUTABLE[_update_hour(last_update_time)]
    return "locked" if int(period_index) < first else "mutable"


def enrich_dispatch(run: DayRun) -> pd.DataFrame:
    raw = run.dispatch.copy()
    n = len(raw)
    if n != T:
        raise ValueError(f"{run.date} dispatch has {n} rows, expected {T}")
    soc_end = raw["soc_kwh"].to_numpy(float)
    soc_start = np.empty(n, dtype=float)
    soc_start[0] = float(run.summary["soc_start_kwh"])
    soc_start[1:] = soc_end[:-1]
    periods = raw["period"].to_numpy(int) if "period" in raw.columns else np.arange(n)
    last_update = raw["last_update_time"].astype(str)
    frame = pd.DataFrame(
        {
            "date": run.date,
            "period_index": periods,
            "time_label": raw["time"].astype(str),
            "g0_kwh": raw["planned_g0_kwh"].to_numpy(float),
            "g_final_kwh": raw["final_g_kwh"].to_numpy(float),
            "x_kwh": raw["normal_x_kwh"].to_numpy(float),
            "emergency_kwh": raw["emergency_kwh"].to_numpy(float),
            "charge_kwh": raw["charge_kwh"].to_numpy(float),
            "discharge_kwh": raw["discharge_kwh"].to_numpy(float),
            "curtailment_kwh": raw["curtailment_kwh"].to_numpy(float),
            "soc_start_kwh": soc_start,
            "soc_end_kwh": soc_end,
            "actual_load_kwh": raw["load_kwh"].to_numpy(float),
            "actual_pv_kwh": raw["actual_pv_kwh"].to_numpy(float),
            "locked_or_mutable": [
                locked_or_mutable_label(int(t), u) for t, u in zip(periods, last_update)
            ],
            "last_update_time": last_update,
            "price": raw["price"].to_numpy(float),
            "phi_yuan": raw["phi_yuan"].to_numpy(float),
            "emergency_cost_yuan": raw["emergency_cost_yuan"].to_numpy(float),
            "run_id": run.run_id,
            "year_end_boundary": run.summary.get("year_end_boundary", YEAR_END_BOUNDARY_A),
            "year_end_soc_kwh": YEAR_END_SOC_A_KWH,
            "strategy": run.strategy,
            "balance_residual_kwh": raw["balance_residual_kwh"].to_numpy(float)
            if "balance_residual_kwh" in raw.columns
            else np.zeros(n),
        }
    )
    return frame[TRAJECTORY_WRITE_COLS]


def has_required_trajectory(frame: pd.DataFrame) -> bool:
    return all(col in frame.columns for col in REQUIRED_TRAJECTORY_COLS) and len(frame) == T


def write_daily_trajectory(frame: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    missing = [col for col in REQUIRED_TRAJECTORY_COLS if col not in frame.columns]
    if missing:
        raise ValueError(f"{path} missing trajectory columns: {missing}")
    if len(frame) != T:
        raise ValueError(f"{path} has {len(frame)} rows, expected {T}")
    frame.to_csv(path, index=False)


def load_daily_trajectory(path: Path) -> pd.DataFrame:
    frame = pd.read_csv(path)
    if not has_required_trajectory(frame):
        raise ValueError(f"{path} is not a streamed 10-minute export trajectory")
    return frame.sort_values("period_index").reset_index(drop=True)


def summarize_trajectory(frame: pd.DataFrame) -> dict:
    return {
        "date": str(frame["date"].iloc[0]),
        "strategy": RESULT3_OFFICIAL_STRATEGY,
        "year_end_boundary": YEAR_END_BOUNDARY_A,
        "year_end_soc_kwh": YEAR_END_SOC_A_KWH,
        "total_cost_yuan": float(frame["phi_yuan"].sum() + frame["emergency_cost_yuan"].sum()),
        "settlement_cost_yuan": float(frame["phi_yuan"].sum()),
        "emergency_cost_yuan": float(frame["emergency_cost_yuan"].sum()),
        "emergency_kwh": float(frame["emergency_kwh"].sum()),
        "curtailment_kwh": float(frame["curtailment_kwh"].sum()),
        "charge_kwh": float(frame["charge_kwh"].sum()),
        "discharge_kwh": float(frame["discharge_kwh"].sum()),
        "soc_start_kwh": float(frame["soc_start_kwh"].iloc[0]),
        "soc_end_kwh": float(frame["soc_end_kwh"].iloc[-1]),
        "g0_kwh": float(frame["g0_kwh"].sum()),
        "g_final_kwh": float(frame["g_final_kwh"].sum()),
        "planned_energy_cost_yuan": float((frame["price"] * frame["g0_kwh"]).sum()),
        "max_balance_residual_kwh": float(frame["balance_residual_kwh"].abs().max())
        if "balance_residual_kwh" in frame.columns
        else 0.0,
    }


def _excel_date(value: object) -> datetime:
    ts = pd.Timestamp(value)
    return datetime(ts.year, ts.month, ts.day)


def _time_headers(ws) -> list[str]:
    headers = [str(ws.cell(1, COL_FIRST_TIME + t).value) for t in range(N_TIME_COLS)]
    if headers[0] != "0:10-0:20" or headers[-1] != "0:00-0:10+1":
        raise ValueError(f"unexpected result3 time headers: {headers[0]!r} .. {headers[-1]!r}")
    return headers


def _fill_energy_grid(ws, by_date: dict[str, pd.DataFrame], value_col: str, fee_col: str) -> None:
    dates = export_dates()
    if ws.max_row != 1 + RESULT3_EXPORT_N_DAYS:
        raise ValueError(f"{ws.title} expected {1 + RESULT3_EXPORT_N_DAYS} rows, got {ws.max_row}")
    for offset, date in enumerate(dates):
        row = 2 + offset
        cell_date = _excel_date(ws.cell(row, 1).value)
        if cell_date.strftime("%Y-%m-%d") != date:
            raise ValueError(f"{ws.title} row {row} date {cell_date} != {date}")
        frame = by_date[date].sort_values("period_index")
        if len(frame) != T:
            raise ValueError(f"{date} {value_col} has {len(frame)} periods")
        values = frame[value_col].to_numpy(float)
        for t, val in enumerate(values):
            if not np.isfinite(val):
                raise ValueError(f"{ws.title} {date} t={t} is not finite")
            ws.cell(row, COL_FIRST_TIME + t).value = float(val)
        ws.cell(row, COL_DAY_QTY).value = float(np.sum(values))
        ws.cell(row, COL_DAY_FEE).value = float(frame[fee_col].sum())
        ws.cell(row, 1).value = cell_date
        ws.cell(row, 1).number_format = "yyyy-mm-dd"


def _rewrite_charge_sheet(ws, by_date: dict[str, pd.DataFrame]) -> None:
    header = [ws.cell(1, c).value for c in range(1, 7)]
    if header != ["日期", "时间段", "充电量", "放电量", "时刻", "储电量"]:
        raise ValueError(f"unexpected 充放电量 header: {header}")
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    dates = export_dates()
    row = 2
    for date in dates:
        frame = by_date[date].sort_values("period_index")
        charge = frame["charge_kwh"].to_numpy(float)
        discharge = frame["discharge_kwh"].to_numpy(float)
        soc_start = float(frame["soc_start_kwh"].iloc[0])
        soc_end = float(frame["soc_end_kwh"].iloc[-1])
        for b, (label, start, end) in enumerate(FOUR_HOUR_BLOCKS):
            ws.cell(row, 1).value = _excel_date(date) if b == 0 else None
            if b == 0:
                ws.cell(row, 1).number_format = "yyyy-mm-dd"
            ws.cell(row, 2).value = label
            ws.cell(row, 3).value = float(np.sum(charge[start:end]))
            ws.cell(row, 4).value = float(np.sum(discharge[start:end]))
            if b == 0:
                ws.cell(row, 5).value = time(0, 0)
                ws.cell(row, 5).number_format = "h:mm"
                ws.cell(row, 6).value = soc_start
            elif b == 1:
                ws.cell(row, 5).value = "24:00"
                ws.cell(row, 6).value = soc_end
            else:
                ws.cell(row, 5).value = None
                ws.cell(row, 6).value = None
            row += 1
    expected = 1 + RESULT3_EXPORT_N_DAYS * len(FOUR_HOUR_BLOCKS)
    if ws.max_row != expected:
        raise ValueError(f"充放电量 expected {expected} rows, got {ws.max_row}")


def _rewrite_emergency_sheet(ws, by_date: dict[str, pd.DataFrame], interval_headers: list[str]) -> int:
    header = [ws.cell(1, c).value for c in range(1, 4)]
    if header != ["日期", "购电时间段", "购电量"]:
        raise ValueError(f"unexpected 紧急购电量 header: {header}")
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    n_rows = 0
    row = 2
    for date in export_dates():
        frame = by_date[date].sort_values("period_index")
        emergency = frame["emergency_kwh"].to_numpy(float)
        for t, value in enumerate(emergency):
            if value <= EMERGENCY_EPS_KWH:
                continue
            ws.cell(row, 1).value = _excel_date(date)
            ws.cell(row, 1).number_format = "yyyy-mm-dd"
            ws.cell(row, 2).value = interval_headers[t]
            ws.cell(row, 3).value = float(value)
            n_rows += 1
            row += 1
    return n_rows


def _write_notes_sheet(wb, source_commit: str, n_emergency: int, export_cost: float) -> None:
    if NOTES_SHEET in wb.sheetnames:
        del wb[NOTES_SHEET]
    ws = wb.create_sheet(NOTES_SHEET)
    lines = [
        "C 题第 3 问正式导出（唯一主路径）",
        "主策略：M1_M6（0:00 制定 g0；6:00/12:00/18:00 评估是否调整未执行承诺）",
        "年末 SOC：A / A_q2_aligned / 1200 kWh",
        "负荷信息：causal_load_main；光伏映射：linear_anchor_main；调整结算：anchor_final_main",
        "初始 SOC：2025-01-01 00:00 = 6000 kWh，跨日连续 E_{d+1,0}=E_{d,144}",
        "输出区间：2025-02-01 至 2025-12-31（334 天）。1 月只用于预热、预测历史和 SOC 连续，不写入本表。",
        "单位：购电/充放电/紧急电量为 kWh；全天购电费为元。",
        "计划购电量：日初承诺 g0，时间列按附件 00:10 … 0:00+1 原顺序，不旋转。全天购电量为行合计；全天购电费为附件1电价与 g0 之积（不含紧急购电惩罚）。",
        "调整购电量：最终有效承诺 gF；未调整时段等于 g0，无空值。全天购电费为实际结算 φ(g0,gF)+5 p e。",
        "充放电量：按执行轨迹的充电量 c、放电量 r 分 6 个四小时段合计；每日第一行储电量为 0:00 SOC，第二行为 24:00 SOC。SOC 是存量，不可跨行或跨日求和。",
        f"紧急购电量：仅列出 e>0 的时段（本文件 {n_emergency} 行）。不得把可由电池平衡的普通缺口列入。",
        f"数据来源 commit：{source_commit}",
        "本文件不含 B 路径、oracle、中间试验或其它策略。",
        f"2–12 月输出区间实际结算合计 {export_cost:.2f} 元；不可用该数字替代 1–12 月年度主结论 {RESULT3_ANNUAL_COST_YUAN:,.2f} 元。",
    ]
    font = Font(name="宋体", size=11)
    for i, line in enumerate(lines, start=1):
        cell = ws.cell(i, 1)
        cell.value = line
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 22
    ws.column_dimensions["A"].width = 118


def write_result3_xlsx(
    by_date: dict[str, pd.DataFrame],
    dest: Path = RESULT3_XLSX,
    template: Path = RESULT3_TEMPLATE,
    source_commit: str | None = None,
) -> dict:
    dates = export_dates()
    missing = [d for d in dates if d not in by_date]
    if missing:
        raise FileNotFoundError(f"missing export-window trajectories: {missing[:5]}")
    if not template.exists():
        raise FileNotFoundError(f"result3 template not found: {template}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, dest)
    wb = load_workbook(dest)
    for name in (PURCHASE_SHEET, ADJUST_SHEET, CHARGE_SHEET, EMERGENCY_SHEET):
        if name not in wb.sheetnames:
            raise ValueError(f"template missing sheet {name}")
    plan = wb[PURCHASE_SHEET]
    adjust = wb[ADJUST_SHEET]
    interval_headers = _time_headers(plan)
    if _time_headers(adjust) != interval_headers:
        raise ValueError("计划/调整 time headers differ")
    fee_plan = {}
    fee_adjust = {}
    grids = {}
    for date in dates:
        frame = by_date[date].copy()
        frame["plan_fee_yuan"] = frame["price"] * frame["g0_kwh"]
        frame["actual_fee_yuan"] = frame["phi_yuan"] + frame["emergency_cost_yuan"]
        grids[date] = frame
        fee_plan[date] = float(frame["plan_fee_yuan"].sum())
        fee_adjust[date] = float(frame["actual_fee_yuan"].sum())
    _fill_energy_grid(plan, grids, "g0_kwh", "plan_fee_yuan")
    _fill_energy_grid(adjust, grids, "g_final_kwh", "actual_fee_yuan")
    _rewrite_charge_sheet(wb[CHARGE_SHEET], grids)
    n_emergency = _rewrite_emergency_sheet(wb[EMERGENCY_SHEET], grids, interval_headers)
    export_cost = float(sum(fee_adjust.values()))
    _write_notes_sheet(wb, source_commit or git_hash(), n_emergency, export_cost)
    for sheet in (PURCHASE_SHEET, ADJUST_SHEET):
        ws = wb[sheet]
        ws.cell(1, COL_DAY_QTY).value = "全天购电量"
        ws.cell(1, COL_DAY_FEE).value = "全天购电费"
        ws.column_dimensions[get_column_letter(COL_DAY_QTY)].width = 14
        ws.column_dimensions[get_column_letter(COL_DAY_FEE)].width = 14
    wb.save(dest)
    return {
        "path": str(dest),
        "n_export_days": len(dates),
        "n_emergency_rows": n_emergency,
        "export_window_cost_yuan": export_cost,
        "interval_headers_first": interval_headers[0],
        "interval_headers_last": interval_headers[-1],
        "planned_window_cost_yuan": float(sum(fee_plan.values())),
    }


def _blank_or_ref(value: object) -> bool:
    if value is None:
        return True
    text = str(value)
    return "#REF!" in text or text.strip() == ""


def inspect_result3_workbook(path: Path) -> dict:
    wb = load_workbook(path, data_only=False)
    plan = wb[PURCHASE_SHEET]
    adjust = wb[ADJUST_SHEET]
    charge = wb[CHARGE_SHEET]
    emergency = wb[EMERGENCY_SHEET]
    anomalies: list[str] = []
    if NOTES_SHEET not in wb.sheetnames:
        anomalies.append("missing 说明 sheet")
    dates = export_dates()
    if plan.max_row != 1 + RESULT3_EXPORT_N_DAYS or adjust.max_row != 1 + RESULT3_EXPORT_N_DAYS:
        anomalies.append("purchase sheets do not have 334 date rows")
    n_blank_adjust = 0
    n_ref = 0
    plan_dates = []
    for offset, date in enumerate(dates):
        row = 2 + offset
        for sheet in (plan, adjust):
            raw = sheet.cell(row, 1).value
            got = pd.Timestamp(raw).strftime("%Y-%m-%d") if raw is not None else None
            if got != date:
                anomalies.append(f"{sheet.title} row {row} date {got} != {date}")
            for col in range(COL_FIRST_TIME, COL_DAY_FEE + 1):
                val = sheet.cell(row, col).value
                if val is not None and "#REF!" in str(val):
                    n_ref += 1
                if sheet is adjust and col < COL_DAY_QTY and _blank_or_ref(val):
                    n_blank_adjust += 1
                if sheet is plan and col < COL_DAY_QTY and _blank_or_ref(val):
                    anomalies.append(f"{sheet.title} {date} col {col} blank")
                    break
        plan_dates.append(date)
    if n_blank_adjust:
        anomalies.append(f"调整购电量 has {n_blank_adjust} blank interval cells")
    if n_ref:
        anomalies.append(f"workbook contains {n_ref} #REF! cells")
    expected_charge_rows = 1 + RESULT3_EXPORT_N_DAYS * len(FOUR_HOUR_BLOCKS)
    if charge.max_row != expected_charge_rows:
        anomalies.append(f"充放电量 rows {charge.max_row} != {expected_charge_rows}")
    jan_leaks = [
        pd.Timestamp(plan.cell(r, 1).value).month == 1
        for r in range(2, plan.max_row + 1)
        if plan.cell(r, 1).value is not None
    ]
    if any(jan_leaks):
        anomalies.append("January dates leaked into 计划购电量")
    return {
        "n_plan_dates": len(plan_dates),
        "n_adjust_dates": adjust.max_row - 1,
        "n_charge_rows": charge.max_row - 1,
        "n_emergency_rows": max(emergency.max_row - 1, 0),
        "has_notes_sheet": NOTES_SHEET in wb.sheetnames,
        "n_anomalies": len(anomalies),
        "anomalies": anomalies,
        "first_date": plan_dates[0] if plan_dates else None,
        "last_date": plan_dates[-1] if plan_dates else None,
        "n_time_columns": N_TIME_COLS,
    }


def compare_to_annual_summary(
    rerun_daily: pd.DataFrame,
    summary_path: Path,
) -> dict:
    if not summary_path.exists():
        return {"compared": False, "reason": f"missing {summary_path}"}
    archive = pd.read_csv(summary_path)
    archive = archive[
        (archive["strategy"] == RESULT3_OFFICIAL_STRATEGY)
        & (archive["year_end_boundary"] == YEAR_END_BOUNDARY_A)
    ].copy()
    archive["date"] = pd.to_datetime(archive["date"]).dt.strftime("%Y-%m-%d")
    rerun = rerun_daily.copy()
    rerun["date"] = pd.to_datetime(rerun["date"]).dt.strftime("%Y-%m-%d")
    merged = rerun.merge(archive, on="date", suffixes=("_rerun", "_archive"), how="left")
    missing = merged[merged["total_cost_yuan_archive"].isna()]["date"].tolist()
    diffs = []
    fields = [
        ("total_cost_yuan", COST_MATCH_TOL_YUAN),
        ("emergency_kwh", ENERGY_MATCH_TOL_KWH),
        ("curtailment_kwh", ENERGY_MATCH_TOL_KWH),
        ("soc_start_kwh", SOC_MATCH_TOL_KWH),
        ("soc_end_kwh", SOC_MATCH_TOL_KWH),
    ]
    max_abs = {name: 0.0 for name, _tol in fields}
    n_fail = 0
    for _, row in merged.iterrows():
        if pd.isna(row.get("total_cost_yuan_archive")):
            continue
        for name, tol in fields:
            delta = abs(float(row[f"{name}_rerun"]) - float(row[f"{name}_archive"]))
            max_abs[name] = max(max_abs[name], delta)
            if delta > tol:
                n_fail += 1
                diffs.append(
                    {
                        "date": row["date"],
                        "field": name,
                        "rerun": float(row[f"{name}_rerun"]),
                        "archive": float(row[f"{name}_archive"]),
                        "abs_diff": delta,
                    }
                )
    return {
        "compared": True,
        "n_rerun_days": int(len(rerun)),
        "n_archive_days": int(len(archive)),
        "n_unmatched_dates": missing,
        "max_abs_diff": max_abs,
        "n_field_failures": n_fail,
        "failures": diffs[:50],
        "all_match": n_fail == 0 and not missing,
        "tolerance_yuan": COST_MATCH_TOL_YUAN,
        "tolerance_kwh": ENERGY_MATCH_TOL_KWH,
    }


def write_export_audit(
    payload: dict,
    path: Path = RESULT3_EXPORT_AUDIT,
) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    return path
