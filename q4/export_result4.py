"""Export result4-2.xlsx and result4-3.xlsx from streamed 10-minute dispatch.

February–December only. January is warmup. Does not modify result2.xlsx or result3.xlsx.
"""

from __future__ import annotations

import argparse
import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT_PATH = Path(__file__).resolve().parents[1]
if str(ROOT_PATH) not in sys.path:
    sys.path.insert(0, str(ROOT_PATH))

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from q3.config import RESULT3_TEMPLATE, T
from q3.export_result3 import (
    ADJUST_SHEET,
    CHARGE_SHEET,
    COL_DAY_FEE,
    COL_DAY_QTY,
    EMERGENCY_SHEET,
    NOTES_SHEET,
    PURCHASE_SHEET,
    _fill_energy_grid,
    _rewrite_charge_sheet,
    _rewrite_emergency_sheet,
    _time_headers,
    export_dates,
)
from q4.audit import ledger_from_dispatch_q42, ledger_from_dispatch_q43
from q4.config import (
    OUTPUT_DIR,
    Q4_2_DISPATCH_DIR,
    Q4_3_DISPATCH_DIR,
    RESULT4_2_XLSX,
    RESULT4_3_XLSX,
    ROOT,
)

RESULT2_LAYOUT = ROOT / "output" / "result2.xlsx"
RESULT4_2_AUDIT = ROOT / "output" / "result4-2_export_audit.json"
RESULT4_3_AUDIT = ROOT / "output" / "result4-3_export_audit.json"


def git_hash() -> str:
    try:
        return subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def _load_dispatch(path: Path, q4_3: bool) -> pd.DataFrame:
    raw = pd.read_csv(path)
    if len(raw) != T:
        raise ValueError(f"{path} has {len(raw)} rows, expected {T}")
    g0 = raw["q_or_g0_kwh"].to_numpy(float)
    gf = raw["g_final_kwh"].to_numpy(float) if q4_3 else g0
    frame = pd.DataFrame(
        {
            "date": raw["date"].astype(str),
            "period_index": raw["period_index"].to_numpy(int),
            "time_label": raw["time_label"].astype(str),
            "g0_kwh": g0,
            "g_final_kwh": gf,
            "x_kwh": raw["x_kwh"].to_numpy(float),
            "emergency_kwh": raw["emergency_kwh"].to_numpy(float),
            "charge_kwh": raw["charge_kwh"].to_numpy(float),
            "discharge_kwh": raw["discharge_kwh"].to_numpy(float),
            "soc_start_kwh": raw["soc_start_kwh"].to_numpy(float),
            "soc_end_kwh": raw["soc_end_kwh"].to_numpy(float),
            "price": raw["actual_price"].to_numpy(float),
            "balance_residual_kwh": raw["balance_residual_kwh"].to_numpy(float),
        }
    )
    billed = ledger_from_dispatch_q43(raw) if q4_3 else ledger_from_dispatch_q42(raw)
    frame["phi_yuan"] = billed["normal_cost_yuan"] + billed["adjustment_cost_yuan"]
    frame["emergency_cost_yuan"] = billed["emergency_cost_yuan"]
    frame["plan_fee_yuan"] = frame["price"] * frame["g0_kwh"]
    frame["actual_fee_yuan"] = billed["total_cost_yuan"]
    return frame.sort_values("period_index").reset_index(drop=True)


def load_export_window(dispatch_dir: Path, q4_3: bool) -> dict[str, pd.DataFrame]:
    dates = export_dates()
    by_date = {}
    missing = []
    for date in dates:
        path = dispatch_dir / f"dispatch_{date}.csv"
        if not path.exists():
            missing.append(date)
            continue
        by_date[date] = _load_dispatch(path, q4_3)
    if missing:
        raise FileNotFoundError(
            f"{dispatch_dir} missing {len(missing)} export-window days, first {missing[:5]}"
        )
    return by_date


def _write_notes(wb, title: str, lines: list[str]) -> None:
    if NOTES_SHEET in wb.sheetnames:
        del wb[NOTES_SHEET]
    ws = wb.create_sheet(NOTES_SHEET)
    font = Font(name="宋体", size=11)
    for i, line in enumerate([title, *lines], start=1):
        cell = ws.cell(i, 1)
        cell.value = line
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 22
    ws.column_dimensions["A"].width = 118


def _label_fee_columns(wb, sheets: tuple[str, ...]) -> None:
    for name in sheets:
        ws = wb[name]
        ws.cell(1, COL_DAY_QTY).value = "全天购电量"
        ws.cell(1, COL_DAY_FEE).value = "全天购电费"
        ws.column_dimensions[get_column_letter(COL_DAY_QTY)].width = 14
        ws.column_dimensions[get_column_letter(COL_DAY_FEE)].width = 14


def write_result4_2(by_date: dict[str, pd.DataFrame], dest: Path = RESULT4_2_XLSX) -> dict:
    if not RESULT2_LAYOUT.exists():
        raise FileNotFoundError(f"result2 layout not found: {RESULT2_LAYOUT}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(RESULT2_LAYOUT, dest)
    wb = load_workbook(dest)
    plan = wb[PURCHASE_SHEET]
    headers = _time_headers(plan)
    _fill_energy_grid(plan, by_date, "g0_kwh", "plan_fee_yuan")
    _rewrite_charge_sheet(wb[CHARGE_SHEET], by_date)
    n_emergency = _rewrite_emergency_sheet(wb[EMERGENCY_SHEET], by_date, headers)
    export_cost = float(sum(float(frame["actual_fee_yuan"].sum()) for frame in by_date.values()))
    annual = float(pd.read_csv(OUTPUT_DIR / "q4_2_warmup_daily.csv")["total_cost_yuan"].sum())
    _write_notes(
        wb,
        "C 题第 4 问 Q4-2 正式导出（日前 q + 因果 MPC）",
        [
            "额度 q 全日锁定；路径内充放电只用于评估 q，不直接执行。",
            "负荷/光伏：Q2 因果口径；价格：已验收 P1 模块；结算：附件4 实际电价 p*q+5*p*e。",
            "年末 SOC：继承 Q2 已验收规则（最后一日无硬约束）。load_information_case=causal_load_main。",
            "输出区间：2025-02-01 至 2025-12-31（334 天）。1 月只用于预热与 SOC 连续，不写入本表。",
            "单位：kWh 与元。时间列按附件 00:10 … 0:00+1 原顺序。",
            f"紧急购电量仅列出 e>0 的时段（本文件 {n_emergency} 行）。",
            f"数据来源 commit：{git_hash()}",
            f"2–12 月实际结算合计 {export_cost:.2f} 元；1–12 月年度合计 {annual:.2f} 元，二者不可混用。",
            "本文件不含 Q4-3、oracle 或未验收价格。未改写 result2.xlsx / result3.xlsx。",
        ],
    )
    _label_fee_columns(wb, (PURCHASE_SHEET,))
    if ADJUST_SHEET in wb.sheetnames:
        del wb[ADJUST_SHEET]
    wb.save(dest)
    return {
        "path": str(dest),
        "n_export_days": len(by_date),
        "n_emergency_rows": n_emergency,
        "export_window_cost_yuan": export_cost,
        "annual_cost_yuan": annual,
    }


def write_result4_3(by_date: dict[str, pd.DataFrame], dest: Path = RESULT4_3_XLSX) -> dict:
    if not RESULT3_TEMPLATE.exists():
        raise FileNotFoundError(f"result3 template not found: {RESULT3_TEMPLATE}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(RESULT3_TEMPLATE, dest)
    wb = load_workbook(dest)
    plan = wb[PURCHASE_SHEET]
    adjust = wb[ADJUST_SHEET]
    headers = _time_headers(plan)
    _fill_energy_grid(plan, by_date, "g0_kwh", "plan_fee_yuan")
    _fill_energy_grid(adjust, by_date, "g_final_kwh", "actual_fee_yuan")
    _rewrite_charge_sheet(wb[CHARGE_SHEET], by_date)
    n_emergency = _rewrite_emergency_sheet(wb[EMERGENCY_SHEET], by_date, headers)
    export_cost = float(sum(float(frame["actual_fee_yuan"].sum()) for frame in by_date.values()))
    annual = float(pd.read_csv(OUTPUT_DIR / "q4_3_warmup_daily.csv")["total_cost_yuan"].sum())
    _write_notes(
        wb,
        "C 题第 4 问 Q4-3 正式导出（固定 M1_M6）",
        [
            "策略固定为 M1_M6：0:00 制定 g0；6:00/12:00/18:00 只改未执行承诺。",
            "预测层使用因果价格；账本按交付时实际电价 φ=p*gF+0.5*p*|gF-g0|，另加 5*p*e。",
            "年末 SOC：A / A_q2_aligned / 1200 kWh。load_information_case=causal_load_main。",
            "输出区间：2025-02-01 至 2025-12-31（334 天）。1 月只用于预热与 SOC 连续，不写入本表。",
            "计划购电量：g0。调整购电量：gF；未调整时段等于 g0。",
            f"紧急购电量仅列出 e>0 的时段（本文件 {n_emergency} 行）。",
            f"数据来源 commit：{git_hash()}",
            f"2–12 月实际结算合计 {export_cost:.2f} 元；1–12 月年度合计 {annual:.2f} 元，二者不可混用。",
            "本文件不含其它策略或 oracle。未改写 result3.xlsx。",
        ],
    )
    _label_fee_columns(wb, (PURCHASE_SHEET, ADJUST_SHEET))
    wb.save(dest)
    return {
        "path": str(dest),
        "n_export_days": len(by_date),
        "n_emergency_rows": n_emergency,
        "export_window_cost_yuan": export_cost,
        "annual_cost_yuan": annual,
    }


def export_q4_2() -> dict:
    q42 = load_export_window(Q4_2_DISPATCH_DIR, q4_3=False)
    r2 = write_result4_2(q42)
    RESULT4_2_AUDIT.write_text(json.dumps(r2, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"result4_2": r2}


def export_q4_3() -> dict:
    q43 = load_export_window(Q4_3_DISPATCH_DIR, q4_3=True)
    r3 = write_result4_3(q43)
    RESULT4_3_AUDIT.write_text(json.dumps(r3, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return {"result4_3": r3}


def export_all() -> dict:
    out = {}
    out.update(export_q4_2())
    out.update(export_q4_3())
    return out


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export result4 workbooks")
    parser.add_argument("--system", choices=("q4_2", "q4_3", "all"), default="all")
    args = parser.parse_args()
    if args.system == "q4_2":
        report = export_q4_2()
    elif args.system == "q4_3":
        report = export_q4_3()
    else:
        report = export_all()
    print(json.dumps(report, ensure_ascii=False, indent=2))
