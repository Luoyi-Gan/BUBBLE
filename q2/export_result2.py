"""Export a candidate result2.xlsx for the policy-consistent Q2 path.

Writes only under output/q2_policy_consistent/. Never overwrites the signed-off
output/result2.xlsx.
"""

from __future__ import annotations

import hashlib
import json
import shutil
from copy import copy
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from q2.config import (
    CANDIDATE_RESULT2,
    E_MAX_KWH,
    E_MIN_KWH,
    EMERGENCY_PRICE_MULTIPLIER,
    NUMERIC_TOL,
    OFFICIAL_OUTPUT_START,
    POLICY_CONSISTENT_OUTPUT_DIR,
    POWER_LIMIT_KWH,
    RESULT2_SIGNED_OFF_BACKUP,
    RESULT2_TEMPLATE,
    ROOT,
    SIGNED_OFF_RESULT2,
    SIGNED_OFF_RESULT2_SHA256,
    SIMULTANEOUS_CD_TOL,
    T,
)
from q2.data import Q2Data
from q2.pilot import planned_q_hash

PURCHASE_SHEET = "计划购电量"
CHARGE_SHEET = "充放电量"
EMERGENCY_SHEET = "紧急购电量"
OUTPUT_END = "2025-12-31"
N_OUTPUT_DAYS = 334
BLOCKS_PER_DAY = 6
LAST_PERIOD_COL = 1 + T
Q_SUM_COL = 2 + T
COST_COL = 3 + T
FOUR_HOUR_BLOCKS = (
    ("0:00-4:00", 0, 24),
    ("4:00-8:00", 24, 48),
    ("8:00-12:00", 48, 72),
    ("12:00-16:00", 72, 96),
    ("16:00-20:00", 96, 120),
    ("20:00-24:00", 120, 144),
)


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def official_output_dates() -> pd.DatetimeIndex:
    return pd.date_range(OFFICIAL_OUTPUT_START, OUTPUT_END, freq="D")


def minutes_to_clock(minutes: int) -> str:
    if minutes == 24 * 60:
        return "24:00"
    hours, mins = divmod(int(minutes), 60)
    return f"{hours:02d}:{mins:02d}"


def format_emergency_interval(t0: int, t1: int) -> str:
    if not (0 <= t0 <= t1 < T):
        raise ValueError(f"invalid emergency interval [{t0}, {t1}]")
    return f"{minutes_to_clock(t0 * 10)}-{minutes_to_clock((t1 + 1) * 10)}"


def emergency_segments(
    emergency: np.ndarray,
    tol: float = 0.0,
) -> list[tuple[int, int, float]]:
    values = np.asarray(emergency, dtype=float)
    if values.shape != (T,):
        raise ValueError(f"emergency vector must have length {T}")
    segments: list[tuple[int, int, float]] = []
    t = 0
    while t < T:
        if values[t] > tol:
            t0 = t
            total = 0.0
            while t < T and values[t] > tol:
                total += float(values[t])
                t += 1
            segments.append((t0, t - 1, total))
        else:
            t += 1
    return segments


def block_energy(values: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    return np.array([float(np.sum(values[start:end])) for _name, start, end in FOUR_HOUR_BLOCKS])


def assert_isolated_candidate_path(dest: Path) -> None:
    dest_resolved = dest.resolve()
    signed = SIGNED_OFF_RESULT2.resolve()
    root = POLICY_CONSISTENT_OUTPUT_DIR.resolve()
    if dest_resolved == signed:
        raise RuntimeError("refusing to overwrite signed-off output/result2.xlsx")
    if dest_resolved.parent != root:
        raise RuntimeError("candidate result2.xlsx must live under output/q2_policy_consistent/")


def preserve_signed_off_backup() -> str:
    if not SIGNED_OFF_RESULT2.exists():
        raise FileNotFoundError(f"missing signed-off {SIGNED_OFF_RESULT2}")
    signed_hash = file_sha256(SIGNED_OFF_RESULT2)
    if signed_hash != SIGNED_OFF_RESULT2_SHA256:
        raise RuntimeError(
            "signed-off output/result2.xlsx hash mismatch; refusing to copy a mutated file"
        )
    RESULT2_SIGNED_OFF_BACKUP.parent.mkdir(parents=True, exist_ok=True)
    if RESULT2_SIGNED_OFF_BACKUP.exists():
        backup_hash = file_sha256(RESULT2_SIGNED_OFF_BACKUP)
        if backup_hash != signed_hash:
            raise RuntimeError("existing result2 backup does not match the signed-off file")
    else:
        shutil.copy2(SIGNED_OFF_RESULT2, RESULT2_SIGNED_OFF_BACKUP)
        backup_hash = file_sha256(RESULT2_SIGNED_OFF_BACKUP)
        if backup_hash != signed_hash:
            raise RuntimeError("failed to copy a bit-identical signed-off result2 backup")
    return signed_hash


def _copy_style(source, dest) -> None:
    if source.has_style:
        dest.font = copy(source.font)
        dest.border = copy(source.border)
        dest.fill = copy(source.fill)
        dest.number_format = source.number_format
        dest.alignment = copy(source.alignment)
        dest.protection = copy(source.protection)


def _set_number(cell, value: float | None) -> None:
    cell.value = None if value is None else float(value)
    cell.number_format = "General"


def _sheet_date(value: object) -> pd.Timestamp:
    return pd.Timestamp(value).normalize()


def _load_day_dispatch(dispatch_dir: Path, date_text: str) -> pd.DataFrame:
    path = dispatch_dir / f"dispatch_{date_text}.csv"
    if not path.exists():
        raise FileNotFoundError(f"missing dispatch {path}")
    frame = pd.read_csv(path)
    if len(frame) != T:
        raise AssertionError(f"{date_text} dispatch has {len(frame)} rows, expected {T}")
    if str(frame["time"].iloc[0]) != "00:10" or str(frame["time"].iloc[-1]) != "0:00+1":
        raise AssertionError(f"{date_text} dispatch is not in raw attachment time order")
    return frame


def write_candidate_result2(
    dest: Path,
    dispatch_dir: Path,
    daily: pd.DataFrame,
    data: Q2Data,
    template: Path = RESULT2_TEMPLATE,
) -> None:
    assert_isolated_candidate_path(dest)
    if not template.exists():
        raise FileNotFoundError(f"missing result2 template {template}")
    dates = official_output_dates()
    ledger = daily.copy()
    ledger["date"] = pd.to_datetime(ledger["date"]).dt.strftime("%Y-%m-%d")
    ledger = ledger.loc[ledger["date"].isin(dates.strftime("%Y-%m-%d"))].sort_values("date")
    if list(ledger["date"]) != list(dates.strftime("%Y-%m-%d")):
        raise AssertionError("Feb-Dec ledger dates do not match the 334 official output days")

    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, dest)
    workbook = load_workbook(dest)
    purchase = workbook[PURCHASE_SHEET]
    charge = workbook[CHARGE_SHEET]
    emergency = workbook[EMERGENCY_SHEET]
    if purchase.max_row != 1 + N_OUTPUT_DAYS:
        raise AssertionError("purchase template does not contain 334 dated rows")
    if purchase.cell(1, Q_SUM_COL).value != "全天购电量" or purchase.cell(1, COST_COL).value != "全天购电费":
        raise AssertionError("purchase template totals headers mismatch")
    if str(purchase.cell(1, 2).value) != "0:10-0:20":
        raise AssertionError("purchase template is not in raw 144-period order")
    if str(purchase.cell(1, LAST_PERIOD_COL).value) != "0:00-0:10+1":
        raise AssertionError("purchase template last period header mismatch")

    charge_style_rows = [2 + offset for offset in range(BLOCKS_PER_DAY)]
    emergency_styles = [emergency.cell(2, column) for column in range(1, 4)]
    emergency_row = 2
    for day_index, date in enumerate(dates):
        date_text = date.strftime("%Y-%m-%d")
        row = ledger.loc[ledger["date"] == date_text].iloc[0]
        frame = _load_day_dispatch(dispatch_dir, date_text)
        q = frame["planned_q_kwh"].to_numpy(dtype=float)
        planned_cost = float(data.price @ q)
        purchase_row = 2 + day_index
        sheet_date = _sheet_date(purchase.cell(purchase_row, 1).value)
        if sheet_date != date:
            raise AssertionError(f"purchase date mismatch at row {purchase_row}: {sheet_date} vs {date}")
        for period in range(T):
            _set_number(purchase.cell(purchase_row, 2 + period), float(q[period]))
        _set_number(purchase.cell(purchase_row, Q_SUM_COL), float(np.sum(q)))
        _set_number(purchase.cell(purchase_row, COST_COL), planned_cost)

        charge_base = 2 + day_index * BLOCKS_PER_DAY
        charge_kwh = block_energy(frame["charge_kwh"].to_numpy(dtype=float))
        discharge_kwh = block_energy(frame["discharge_kwh"].to_numpy(dtype=float))
        for offset, (name, _start, _end) in enumerate(FOUR_HOUR_BLOCKS):
            excel_row = charge_base + offset
            style_source_row = charge_style_rows[offset]
            for column in range(1, 7):
                _copy_style(charge.cell(style_source_row, column), charge.cell(excel_row, column))
            charge.cell(excel_row, 1).value = purchase.cell(purchase_row, 1).value if offset == 0 else None
            charge.cell(excel_row, 2).value = name
            _set_number(charge.cell(excel_row, 3), float(charge_kwh[offset]))
            _set_number(charge.cell(excel_row, 4), float(discharge_kwh[offset]))
            if offset == 0:
                charge.cell(excel_row, 5).value = "00:00"
                _set_number(charge.cell(excel_row, 6), float(row["soc_start_kwh"]))
            elif offset == 1:
                charge.cell(excel_row, 5).value = "24:00"
                _set_number(charge.cell(excel_row, 6), float(row["soc_end_kwh"]))
            else:
                charge.cell(excel_row, 5).value = None
                _set_number(charge.cell(excel_row, 6), None)

        segments = emergency_segments(frame["emergency_kwh"].to_numpy(dtype=float))
        if not segments:
            for column in range(1, 4):
                _copy_style(emergency_styles[column - 1], emergency.cell(emergency_row, column))
            emergency.cell(emergency_row, 1).value = purchase.cell(purchase_row, 1).value
            emergency.cell(emergency_row, 2).value = None
            _set_number(emergency.cell(emergency_row, 3), None)
            emergency_row += 1
        else:
            for seg_index, (t0, t1, amount) in enumerate(segments):
                for column in range(1, 4):
                    _copy_style(emergency_styles[column - 1], emergency.cell(emergency_row, column))
                emergency.cell(emergency_row, 1).value = (
                    purchase.cell(purchase_row, 1).value if seg_index == 0 else None
                )
                emergency.cell(emergency_row, 2).value = format_emergency_interval(t0, t1)
                _set_number(emergency.cell(emergency_row, 3), float(amount))
                emergency_row += 1

    leftover_start = emergency_row
    leftover_end = max(emergency.max_row, leftover_start - 1)
    for row_index in range(leftover_start, leftover_end + 1):
        for column in range(1, 4):
            emergency.cell(row_index, column).value = None

    workbook.save(dest)


def _cell_float(sheet: Worksheet, row: int, column: int) -> float:
    value = sheet.cell(row, column).value
    if value in (None, ""):
        return 0.0
    if isinstance(value, datetime):
        raise TypeError(f"{sheet.title}!{sheet.cell(row, column).coordinate} stored a date instead of a number")
    return float(value)


def audit_candidate_result2(
    dest: Path,
    dispatch_dir: Path,
    daily: pd.DataFrame,
    data: Q2Data,
    inherited_soc_kwh: float,
) -> dict:
    assert_isolated_candidate_path(dest)
    workbook = load_workbook(dest, data_only=False)
    purchase = workbook[PURCHASE_SHEET]
    charge = workbook[CHARGE_SHEET]
    emergency = workbook[EMERGENCY_SHEET]
    dates = official_output_dates()
    ledger = daily.copy()
    ledger["date"] = pd.to_datetime(ledger["date"]).dt.strftime("%Y-%m-%d")
    ledger = ledger.loc[ledger["date"].isin(dates.strftime("%Y-%m-%d"))].sort_values("date")
    if len(ledger) != N_OUTPUT_DAYS:
        raise AssertionError("ledger does not cover 334 official output days")

    rows: list[dict] = []
    emergency_row = 2
    previous_soc_end: float | None = None
    max_q_abs = 0.0
    max_cost_abs = 0.0
    max_emergency_abs = 0.0
    max_charge_abs = 0.0
    max_soc_abs = 0.0
    max_balance = 0.0
    max_soc_gap = 0.0
    dated_emergency_rows = 0

    for day_index, date in enumerate(dates):
        date_text = date.strftime("%Y-%m-%d")
        ledger_row = ledger.loc[ledger["date"] == date_text].iloc[0]
        frame = _load_day_dispatch(dispatch_dir, date_text)
        q = frame["planned_q_kwh"].to_numpy(dtype=float)
        x = frame["actual_x_kwh"].to_numpy(dtype=float)
        emergency_kwh = frame["emergency_kwh"].to_numpy(dtype=float)
        charge_kwh = frame["charge_kwh"].to_numpy(dtype=float)
        discharge_kwh = frame["discharge_kwh"].to_numpy(dtype=float)
        purchase_row = 2 + day_index
        sheet_date = _sheet_date(purchase.cell(purchase_row, 1).value)
        if sheet_date != date:
            raise AssertionError(f"xlsx purchase date mismatch: {sheet_date} vs {date}")
        xlsx_q = np.array([_cell_float(purchase, purchase_row, 2 + period) for period in range(T)])
        q_abs = float(np.max(np.abs(xlsx_q - q)))
        xlsx_q_sum = _cell_float(purchase, purchase_row, Q_SUM_COL)
        xlsx_cost = _cell_float(purchase, purchase_row, COST_COL)
        planned_cost = float(data.price @ q)
        cost_abs = abs(xlsx_cost - planned_cost)
        q_sum_abs = abs(xlsx_q_sum - float(np.sum(q)))

        charge_base = 2 + day_index * BLOCKS_PER_DAY
        expected_charge = block_energy(charge_kwh)
        expected_discharge = block_energy(discharge_kwh)
        charge_abs = 0.0
        for offset, (name, _start, _end) in enumerate(FOUR_HOUR_BLOCKS):
            excel_row = charge_base + offset
            date_cell = charge.cell(excel_row, 1).value
            if offset == 0:
                if _sheet_date(date_cell) != date:
                    raise AssertionError(f"charge date mismatch on {date_text}")
            elif date_cell not in (None, ""):
                raise AssertionError(f"charge date repeated on {date_text} block {offset}")
            if charge.cell(excel_row, 2).value != name:
                raise AssertionError(f"charge block label mismatch on {date_text}")
            charge_abs = max(
                charge_abs,
                abs(_cell_float(charge, excel_row, 3) - expected_charge[offset]),
                abs(_cell_float(charge, excel_row, 4) - expected_discharge[offset]),
            )
        start_clock = charge.cell(charge_base, 5).value
        end_clock = charge.cell(charge_base + 1, 5).value
        if start_clock not in ("00:00", datetime.strptime("00:00", "%H:%M").time()):
            raise AssertionError(f"{date_text} 00:00 SOC clock mismatch")
        if end_clock != "24:00":
            raise AssertionError(f"{date_text} 24:00 SOC clock mismatch")
        xlsx_soc_start = _cell_float(charge, charge_base, 6)
        xlsx_soc_end = _cell_float(charge, charge_base + 1, 6)
        soc_start_abs = abs(xlsx_soc_start - float(ledger_row["soc_start_kwh"]))
        soc_end_abs = abs(xlsx_soc_end - float(ledger_row["soc_end_kwh"]))
        last_dispatch_soc = float(frame["soc_kwh"].iloc[-1])
        soc_dispatch_abs = abs(xlsx_soc_end - last_dispatch_soc)
        if previous_soc_end is not None:
            max_soc_gap = max(max_soc_gap, abs(xlsx_soc_start - previous_soc_end))
        previous_soc_end = xlsx_soc_end

        date_cell = emergency.cell(emergency_row, 1).value
        if date_cell in (None, ""):
            raise AssertionError(f"missing dated emergency row for {date_text}")
        if _sheet_date(date_cell) != date:
            raise AssertionError(f"emergency date mismatch: {date_cell} vs {date_text}")
        dated_emergency_rows += 1
        xlsx_emergency = 0.0
        first = True
        while emergency_row <= emergency.max_row:
            current_date = emergency.cell(emergency_row, 1).value
            interval = emergency.cell(emergency_row, 2).value
            amount = emergency.cell(emergency_row, 3).value
            if not first and current_date not in (None, ""):
                break
            if interval in (None, "") and amount in (None, ""):
                if first:
                    emergency_row += 1
                break
            xlsx_emergency += float(amount)
            first = False
            emergency_row += 1
        if float(emergency_kwh.min()) < -NUMERIC_TOL:
            raise AssertionError(f"{date_text} has negative emergency energy")
        e_abs = abs(xlsx_emergency - float(np.sum(np.clip(emergency_kwh, 0.0, None))))
        e_ledger_abs = abs(xlsx_emergency - float(ledger_row["emergency_kwh"]))
        e_recon_abs = e_ledger_abs

        residual = (
            frame["actual_x_kwh"]
            + frame["emergency_kwh"]
            + frame["pv_kwh"]
            - frame["curtailment_kwh"]
            + frame["discharge_kwh"]
            - frame["load_kwh"]
            - frame["charge_kwh"]
        )
        day_index_data = int(data.dates.get_loc(date))
        balance = float(residual.abs().max())
        x_minus_q = float(np.max(x - q))
        cd = float((charge_kwh * discharge_kwh).max())
        if (
            planned_q_hash(q) != ledger_row["planned_q_sha256"]
            or abs(planned_cost - float(ledger_row["planned_cost_yuan"])) > 1e-6
            or abs(float(np.sum(q)) - float(ledger_row["planned_q_kwh"])) > 1e-6
        ):
            raise AssertionError(f"{date_text} dispatch does not match the annual ledger")
        if not np.allclose(frame["load_kwh"].to_numpy(), data.load[day_index_data], atol=1e-12):
            raise AssertionError(f"{date_text} load mismatch versus attachment")
        emergency_cost = float(EMERGENCY_PRICE_MULTIPLIER * data.price @ emergency_kwh)
        if abs(emergency_cost - float(ledger_row["emergency_cost_yuan"])) > 1e-4:
            raise AssertionError(f"{date_text} emergency cost mismatch versus ledger")

        rows.append(
            {
                "date": date_text,
                "q_max_abs_kwh": q_abs,
                "q_sum_abs_kwh": q_sum_abs,
                "planned_cost_abs_yuan": cost_abs,
                "charge_block_abs_kwh": charge_abs,
                "soc_start_abs_kwh": soc_start_abs,
                "soc_end_abs_kwh": soc_end_abs,
                "soc_vs_dispatch_abs_kwh": soc_dispatch_abs,
                "emergency_abs_kwh": e_abs,
                "emergency_vs_ledger_abs_kwh": e_ledger_abs,
                "emergency_recon_abs_kwh": e_recon_abs,
                "max_balance_residual_kwh": balance,
                "max_x_minus_q_kwh": x_minus_q,
                "max_simultaneous_cd_kwh2": cd,
                "charge_power_ok": float(charge_kwh.max()) <= POWER_LIMIT_KWH + NUMERIC_TOL,
                "discharge_power_ok": float(discharge_kwh.max()) <= POWER_LIMIT_KWH + NUMERIC_TOL,
                "soc_bounds_ok": (
                    float(frame["soc_kwh"].min()) >= E_MIN_KWH - NUMERIC_TOL
                    and float(frame["soc_kwh"].max()) <= E_MAX_KWH + NUMERIC_TOL
                ),
            }
        )
        max_q_abs = max(max_q_abs, q_abs, q_sum_abs)
        max_cost_abs = max(max_cost_abs, cost_abs)
        max_emergency_abs = max(max_emergency_abs, e_recon_abs)
        max_charge_abs = max(max_charge_abs, charge_abs)
        max_soc_abs = max(max_soc_abs, soc_start_abs, soc_end_abs, soc_dispatch_abs)
        max_balance = max(max_balance, balance)

    if dated_emergency_rows != N_OUTPUT_DAYS:
        raise AssertionError("emergency sheet does not have one dated row per official day")
    feb1_soc = float(ledger.iloc[0]["soc_start_kwh"])
    if abs(feb1_soc - inherited_soc_kwh) > NUMERIC_TOL:
        raise AssertionError("Feb 1 xlsx/ledger SOC does not inherit January warmup")

    xlsx_planned_q = sum(_cell_float(purchase, row, Q_SUM_COL) for row in range(2, 2 + N_OUTPUT_DAYS))
    xlsx_planned_cost = sum(_cell_float(purchase, row, COST_COL) for row in range(2, 2 + N_OUTPUT_DAYS))
    xlsx_emergency = 0.0
    for row in range(2, emergency.max_row + 1):
        amount = emergency.cell(row, 3).value
        if amount not in (None, ""):
            xlsx_emergency += float(amount)

    checks = {
        "n_days": N_OUTPUT_DAYS,
        "n_periods": T,
        "units": {"energy": "kWh", "cost": "yuan"},
        "purchase_sheet_rows": 1 + N_OUTPUT_DAYS,
        "charge_sheet_rows": 1 + N_OUTPUT_DAYS * BLOCKS_PER_DAY,
        "dated_emergency_rows": dated_emergency_rows,
        "feb1_inherited_soc_kwh": inherited_soc_kwh,
        "feb1_xlsx_soc_kwh": float(ledger.iloc[0]["soc_start_kwh"]),
        "max_q_abs_kwh": max_q_abs,
        "max_planned_cost_abs_yuan": max_cost_abs,
        "max_charge_block_abs_kwh": max_charge_abs,
        "max_soc_abs_kwh": max_soc_abs,
        "max_soc_continuity_gap_kwh": max_soc_gap,
        "max_emergency_recon_abs_kwh": max_emergency_abs,
        "max_balance_residual_kwh": max_balance,
        "xlsx_planned_q_kwh": xlsx_planned_q,
        "ledger_planned_q_kwh": float(ledger["planned_q_kwh"].sum()),
        "xlsx_planned_cost_yuan": xlsx_planned_cost,
        "ledger_planned_cost_yuan": float(ledger["planned_cost_yuan"].sum()),
        "xlsx_emergency_kwh": xlsx_emergency,
        "ledger_emergency_kwh": float(ledger["emergency_kwh"].sum()),
        "ledger_total_cost_yuan": float(ledger["total_cost_yuan"].sum()),
        "pass": (
            max_q_abs < 1e-8
            and max_cost_abs < 1e-6
            and max_charge_abs < 1e-8
            and max_soc_abs < 1e-8
            and max_soc_gap < NUMERIC_TOL
            and max_emergency_abs < 1e-8
            and max_balance < NUMERIC_TOL
            and abs(xlsx_planned_q - float(ledger["planned_q_kwh"].sum())) < 1e-6
            and abs(xlsx_planned_cost - float(ledger["planned_cost_yuan"].sum())) < 1e-4
            and abs(xlsx_emergency - float(ledger["emergency_kwh"].sum())) < 1e-4
            and all(row["charge_power_ok"] and row["discharge_power_ok"] and row["soc_bounds_ok"] for row in rows)
            and all(row["max_x_minus_q_kwh"] < NUMERIC_TOL for row in rows)
            and all(row["max_simultaneous_cd_kwh2"] <= SIMULTANEOUS_CD_TOL for row in rows)
            and dest.resolve() != SIGNED_OFF_RESULT2.resolve()
            and file_sha256(SIGNED_OFF_RESULT2) == SIGNED_OFF_RESULT2_SHA256
        ),
    }
    return {"summary": checks, "daily": rows}


def write_result2_audit(audit: dict, output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(audit["daily"]).to_csv(output / "result2_cell_audit.csv", index=False)
    payload = {
        "candidate_result2": str(CANDIDATE_RESULT2.relative_to(ROOT)),
        "signed_off_result2_unchanged": True,
        "signed_off_sha256": SIGNED_OFF_RESULT2_SHA256,
        "backup": str(RESULT2_SIGNED_OFF_BACKUP.relative_to(ROOT)),
        **audit["summary"],
    }
    (output / "result2_cell_audit.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
