from __future__ import annotations

from pathlib import Path
import shutil

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

from q1.config import (
    ALLOWED_CHARGE_CELLS,
    ALLOWED_PURCHASE_CELLS,
    CHARGE_SHEET,
    DELTA_H,
    EPS_C,
    ETA_RT,
    FIG_DIR,
    FOUR_HOUR_BLOCKS,
    OFFICIAL_MODEL,
    PURCHASE_SHEET,
    ROW_ORDER_POLICY,
    SIMUL_CD_TOL,
    T,
    TABLE2_CHARGE_VAR,
    TABLE2_DISCHARGE_VAR,
    TABLE2_ENERGY_SIDE,
)
from q1.optimize import SolveResult


def _cjk_rc() -> None:
    plt.rcParams["font.sans-serif"] = ["PingFang SC", "Heiti SC", "Songti SC", "Arial Unicode MS", "DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False


def write_plan_csv(path: Path, labels: list[str], data: pd.DataFrame, res: SolveResult) -> None:
    balance = res.g + data["pv_kwh"].to_numpy() + res.d - data["load_kwh"].to_numpy() - res.c - res.s
    plan = pd.DataFrame(
        {
            "time": labels,
            "price": data["price"],
            "load_kwh": data["load_kwh"],
            "pv_kwh": data["pv_kwh"],
            "grid_purchase_kwh": res.g,
            "charge_kwh": res.c,
            "discharge_kwh": res.d,
            "soc_kwh": res.E[1:],
            "curtailment_kwh": res.s,
            "balance_residual_kwh": balance,
        }
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    plan.to_csv(path, index=False)


def write_summary_csv(path: Path, res: SolveResult) -> None:
    p = res.params
    rows = [
        ("official_model", OFFICIAL_MODEL, "1", "正式方案；M2–M4 仅对照"),
        ("grid_purchase_kwh", res.grid_purchase, "kWh", "全天计划购电量合计"),
        ("purchase_cost_yuan", res.purchase_cost, "yuan", "全天购电费"),
        ("peak_grid_kw", res.peak_grid_kw, "kW", "购电功率峰值 max(g_t)/Δ"),
        ("curtailment_kwh", res.curtailment, "kWh", "全天弃光量"),
        ("throughput_kwh", res.throughput, "kWh", "充放电吞吐量 Q=sum(c+d)"),
        ("soc_min_kwh", res.soc_min, "kWh", "SOC 最小值含 E0"),
        ("soc_max_kwh", res.soc_max, "kWh", "SOC 最大值含 E0"),
        ("E0_kwh", res.E[0], "kWh", "0:00 储电量"),
        ("E144_kwh", res.E[-1], "kWh", "24:00 储电量"),
        ("max_balance_residual_kwh", res.max_balance_residual, "kWh", "能量平衡残差绝对值最大"),
        ("max_simultaneous_charge_discharge", res.max_cd, "kWh^2", "max(c_t * d_t)"),
        ("max_s_minus_pv_kwh", res.max_s_minus_pv, "kWh", "max(s_t-P_t)，应≤0"),
        ("eta_rt", ETA_RT, "1", "主方案往返效率"),
        ("eta_c", p.eta_c, "1", "充电效率"),
        ("eta_d", p.eta_d, "1", "放电效率"),
        ("c_max_kwh", p.c_max_kwh, "kWh", "每时段充电上限 power_limit_kw*Δ"),
        ("d_max_kwh", p.d_max_kwh, "kWh", "每时段放电上限"),
        ("power_limit_kw", p.power_limit_kw, "kW", "充放电功率上限"),
        ("grid_limit_kw", "" if p.grid_limit_kw is None else p.grid_limit_kw, "kW", "主方案无上限"),
        ("eps_C", EPS_C, "yuan", "词典序成本容差（仅对照）"),
        ("table2_energy_side", TABLE2_ENERGY_SIDE, "1", "表2计量侧，待队长最终签收"),
        ("table2_charge_var", TABLE2_CHARGE_VAR, "1", "充电量汇总变量"),
        ("table2_discharge_var", TABLE2_DISCHARGE_VAR, "1", "放电量汇总变量"),
        ("row_order_policy", ROW_ORDER_POLICY, "1", "附件原始行序口径（队长确认）"),
    ]
    pd.DataFrame(rows, columns=["metric", "value", "unit", "meaning"]).to_csv(path, index=False)


def write_comparison_csv(path: Path, results: list[SolveResult]) -> None:
    rows = []
    for r in results:
        rows.append(
            {
                "case": r.name,
                "solver": "HIGHS",
                "purchase_cost_yuan": r.purchase_cost,
                "grid_purchase_kwh": r.grid_purchase,
                "peak_grid_kw": r.peak_grid_kw,
                "curtailment_kwh": r.curtailment,
                "throughput_kwh": r.throughput,
                "max_simultaneous_charge_discharge": r.max_cd,
                "soc_min_kwh": r.soc_min,
                "soc_max_kwh": r.soc_max,
                "max_balance_residual_kwh": r.max_balance_residual,
                "max_s_minus_pv_kwh": r.max_s_minus_pv,
                "solve_seconds": r.solve_seconds,
                "pass": r.passed,
            }
        )
    pd.DataFrame(rows).to_csv(path, index=False)


def write_baseline_csv(
    path: Path,
    g_base: np.ndarray,
    s_base: np.ndarray,
    price: np.ndarray,
    m1: SolveResult,
) -> None:
    g_sum = float(np.sum(g_base))
    s_sum = float(np.sum(s_base))
    cost = float(price @ g_base)
    save = cost - m1.purchase_cost
    ratio = save / cost if cost else np.nan
    rows = [
        ("baseline_grid_purchase_kwh", g_sum, "kWh", "无储能无售电：g=max(L-P,0)"),
        ("baseline_purchase_cost_yuan", cost, "yuan", "基准购电费"),
        ("baseline_curtailment_kwh", s_sum, "kWh", "基准弃光 s=max(P-L,0)"),
        ("m1_grid_purchase_kwh", m1.grid_purchase, "kWh", "正式 M1 总购电量"),
        ("m1_purchase_cost_yuan", m1.purchase_cost, "yuan", "正式 M1 购电费"),
        ("absolute_cost_saving_yuan", save, "yuan", "基准费用减 M1 费用"),
        ("cost_saving_ratio", ratio, "1", "节省比例（相对基准）"),
    ]
    pd.DataFrame(rows, columns=["metric", "value", "unit", "meaning"]).to_csv(path, index=False)


def write_sensitivity_csv(path: Path, rows: list[dict]) -> None:
    cols = [
        "case",
        "eta_c",
        "eta_d",
        "nominal_capacity_kwh",
        "soc_min_kwh",
        "soc_max_kwh",
        "power_limit_kw",
        "grid_limit_kw",
        "purchase_cost_yuan",
        "grid_purchase_kwh",
        "peak_grid_kw",
        "curtailment_kwh",
        "throughput_kwh",
        "soc_observed_min_kwh",
        "soc_observed_max_kwh",
        "pass",
    ]
    pd.DataFrame(rows)[cols].to_csv(path, index=False)


def sensitivity_row(case: str, res: SolveResult) -> dict:
    p = res.params
    return {
        "case": case,
        "eta_c": p.eta_c,
        "eta_d": p.eta_d,
        "nominal_capacity_kwh": p.nominal_capacity_kwh,
        "soc_min_kwh": p.e_min_kwh,
        "soc_max_kwh": p.e_max_kwh,
        "power_limit_kw": p.power_limit_kw,
        "grid_limit_kw": "" if p.grid_limit_kw is None else p.grid_limit_kw,
        "purchase_cost_yuan": res.purchase_cost,
        "grid_purchase_kwh": res.grid_purchase,
        "peak_grid_kw": res.peak_grid_kw,
        "curtailment_kwh": res.curtailment,
        "throughput_kwh": res.throughput,
        "soc_observed_min_kwh": res.soc_min,
        "soc_observed_max_kwh": res.soc_max,
        "pass": res.passed,
    }


def write_model_selection_md(path: Path, results: dict[str, SolveResult], official: str) -> None:
    m1, m2, m3, m4 = results["M1"], results["M2"], results["M3"], results["M4"]
    lines = [
        "# Q1 模型形式对照与选择",
        "",
        f"- 正式方案为 **{official}**（纯成本最优 LP）。M2–M4 仅作数值对照，不得覆盖 `result1.xlsx` / 正式图。",
        "- M2 是在成本最优容差内再最小化吞吐量的数值择优，**不能**据此声称降低循环寿命损耗或延长电池寿命。",
        f"- 词典序成本容差 `eps_C` = {EPS_C} 元（仅用于 M2/M4 对照）。",
        f"- 效率口径 A（正式）：η_rt = {ETA_RT}，η_c = η_d = sqrt(0.9) ≈ {m1.params.eta_c:.12f}。",
        "",
        "## 结果摘要",
        "",
        "| case | cost | G | peak kW | S | Q | max c*d | SOC min/max | residual | seconds | pass |",
        "|---|---:|---:|---:|---:|---:|---:|---|---:|---:|---|",
    ]
    for key in ("M1", "M2", "M3", "M4"):
        r = results[key]
        lines.append(
            f"| {key} | {r.purchase_cost:.6f} | {r.grid_purchase:.6f} | {r.peak_grid_kw:.6f} | "
            f"{r.curtailment:.6f} | {r.throughput:.6f} | {r.max_cd:.3e} | "
            f"{r.soc_min:.4f}/{r.soc_max:.4f} | {r.max_balance_residual:.3e} | "
            f"{r.solve_seconds:.3f} | {r.passed} |"
        )
    lines += [
        "",
        "## 约束检查",
        "",
        f"- M1 同时充放电 max(c d) = {m1.max_cd:.6e}；阈值 {SIMUL_CD_TOL}。",
        f"- M1 与 M3 成本差 = {abs(m1.purchase_cost - m3.purchase_cost):.6e}。",
        f"- M2 相对 M1 成本增量 = {m2.purchase_cost - m1.purchase_cost:.6e}（应 ≤ eps_C）。",
        f"- M2 相对 M1 吞吐量变化 ΔQ = {m2.throughput - m1.throughput:.6f}。",
        f"- M2 与 M4 成本差 = {abs(m2.purchase_cost - m4.purchase_cost):.6e}，ΔQ = {m2.throughput - m4.throughput:.6f}。",
        f"- M1 max(s-P) = {m1.max_s_minus_pv:.6e}。",
        "",
        "## 推荐",
        "",
        "- 正式采用 M1：目标仅为购电费，与题面一致。",
    ]
    if m1.passed and m3.passed and abs(m1.purchase_cost - m3.purchase_cost) <= EPS_C and m1.max_cd <= SIMUL_CD_TOL:
        lines.append("- M1 与 M3 成本相同且同时充放电仅为数值残差，正式模型不需要二元互斥变量。")
    if m2.passed and m2.purchase_cost <= m1.purchase_cost + EPS_C + 1e-9:
        lines.append("- M2 可在对照表中作为同等费用下的低吞吐量解，但不作为正式答卷。")
    lines += ["", "## 求解状态", ""]
    for key, r in results.items():
        lines.append(f"- {key}: status={r.status}; notes={r.notes or 'ok'}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_validation_md(
    path: Path,
    data: pd.DataFrame,
    labels: list[str],
    official: SolveResult,
    attach_path: Path,
    template_path: Path,
    price_min: float,
    gmax_min_kw: float,
) -> None:
    mapping_rows = []
    for i in [0, 1, 2, T // 2 - 1, T - 3, T - 2, T - 1]:
        mapping_rows.append(
            f"| {i + 1} | {data['attach_time'].iloc[i]} | {labels[i]} |"
        )
    p = official.params
    lines = [
        "# Q1 验收",
        "",
        f"- 附件1：`{attach_path}`",
        f"- 模板：`{template_path}`",
        f"- 正式方案：{official.name}（{OFFICIAL_MODEL}），solver=HIGHS，status={official.status}",
        f"- **附件原始行序口径（队长确认）**：不按时间排序，不把末行旋转到首行。",
        f"- 行序策略代码标记：`{ROW_ORDER_POLICY}`",
        "",
        "## 时间标签映射",
        "",
        "内部模型用整数索引 t=1..144。SOC 独立使用 E_0,...,E_144，其中 E_0=E_144=6000 kWh 对应 0:00 与 24:00。",
        "导出时按官方模板**既有行序**填写第 t 个购电结果，不新增 `0:00-0:10`，不改标签。",
        "",
        "规则：内部索引 t → 附件1第 t 条数据 → 模板第 t 条时间标签。",
        "",
        "| t | 附件1时间 | 模板时间段 |",
        "|---:|---|---|",
        *mapping_rows,
        "",
        "附件1首条为 `00:10:00`，模板首条为 `0:10-0:20`；本题按行序一一对应。",
        "四小时充放电汇总按内部索引每 24 个时段分块：t=1..24, 25..48, …, 121..144。",
        "",
        "## 输入检查",
        "",
        f"- 价格最小值 = {price_min:.6g} 元/kWh。当前附件无负价；若出现负价必须报错，不得沿用“无需同时充放电”的结论。",
        "",
        "## 表 2 计量侧（待队长最终签收）",
        "",
        f"- 集中定义：`TABLE2_ENERGY_SIDE={TABLE2_ENERGY_SIDE}`。",
        f"- 充电量 = `{TABLE2_CHARGE_VAR}`（公共母线侧充入储能）。",
        f"- 放电量 = `{TABLE2_DISCHARGE_VAR}`（公共母线侧储能放电）。",
        "- 未换算电池内部能量；签收前不得在多处隐式改写。",
        "",
        "## 数值验收",
        "",
        f"- 时段数：{T}",
        f"- 功率已乘 Δ={DELTA_H} h 转为 kWh",
        f"- max |能量平衡残差| = {official.max_balance_residual:.6e} kWh（阈值 1e-6）",
        f"- max c,d = {official.c.max():.6f}, {official.d.max():.6f} kWh（上限 {p.c_max_kwh:.9f}）",
        f"- SOC min/max = {official.soc_min:.6f} / {official.soc_max:.6f} kWh（[{p.e_min_kwh}, {p.e_max_kwh}]）",
        f"- E0 = {official.E[0]:.12f}，E144 = {official.E[-1]:.12f}，|E144-E0|={abs(official.E[-1]-p.e0_kwh):.6e}",
        f"- max(c_t d_t) = {official.max_cd:.6e}（阈值 1e-4）",
        f"- max(s_t-P_t) = {official.max_s_minus_pv:.6e}",
        f"- 峰值购电功率 = {official.peak_grid_kw:.6f} kW（主方案无 PCC 上限）",
        f"- PCC 敏感性：最小可行峰值约 {gmax_min_kw:.6f} kW，该值**不是**题给参数。",
        f"- 验收通过：{official.passed}",
        "",
    ]
    if not official.passed:
        lines.append(f"**未通过：** {official.notes}")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def table2_block_energy(res: SolveResult) -> tuple[np.ndarray, np.ndarray]:
    """PCC AC-bus side: charge = c_t, discharge = d_t. Single definition for Table 2."""
    charge = np.array([float(np.sum(res.c[a:b])) for _n, a, b in FOUR_HOUR_BLOCKS])
    discharge = np.array([float(np.sum(res.d[a:b])) for _n, a, b in FOUR_HOUR_BLOCKS])
    return charge, discharge


def write_result1_xlsx(template: Path, dest: Path, res: SolveResult) -> None:
    dest.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template, dest)
    wb = load_workbook(dest)
    ws_g = wb[PURCHASE_SHEET]
    for i in range(T):
        header = ws_g.cell(row=i + 2, column=1).value
        if header is None:
            raise RuntimeError(f"template row {i+2} missing time label")
        ws_g.cell(row=i + 2, column=2).value = float(res.g[i])
    ws_cd = wb[CHARGE_SHEET]
    charge, discharge = table2_block_energy(res)
    for row, val_c, val_d in zip(range(2, 8), charge, discharge):
        ws_cd.cell(row=row, column=2).value = val_c
        ws_cd.cell(row=row, column=3).value = val_d
    ws_cd.cell(row=2, column=5).value = float(res.E[0])
    ws_cd.cell(row=3, column=5).value = float(res.E[-1])
    wb.save(dest)
    verify_xlsx_only_allowed_cells(template, dest)


def verify_xlsx_only_allowed_cells(template: Path, dest: Path) -> None:
    src = load_workbook(template, data_only=False)
    out = load_workbook(dest, data_only=False)
    for name in src.sheetnames:
        ws_s, ws_o = src[name], out[name]
        for row in ws_s.iter_rows(min_row=1, max_row=ws_s.max_row, max_col=ws_s.max_column):
            for cell in row:
                key = (cell.row, cell.column)
                if name == PURCHASE_SHEET and key in ALLOWED_PURCHASE_CELLS:
                    continue
                if name == CHARGE_SHEET and key in ALLOWED_CHARGE_CELLS:
                    continue
                if ws_o.cell(cell.row, cell.column).value != cell.value:
                    raise RuntimeError(
                        f"modified non-fillable cell {name}!{cell.coordinate}: "
                        f"{cell.value!r} -> {ws_o.cell(cell.row, cell.column).value!r}"
                    )


def write_figures(data: pd.DataFrame, res: SolveResult) -> None:
    _cjk_rc()
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    hours = (np.arange(T) + 1) * DELTA_H

    fig, ax = plt.subplots(figsize=(10, 4.2))
    ax.plot(hours, data["load_kwh"], label="负荷 (kWh/10min)", lw=1.2)
    ax.plot(hours, data["pv_kwh"], label="光伏 (kWh/10min)", lw=1.2)
    ax.plot(hours, res.g, label="购电 (kWh)", lw=1.2)
    ax.plot(hours, res.c, label="充电 (kWh)", lw=1.0)
    ax.plot(hours, res.d, label="放电 (kWh)", lw=1.0)
    ax.set_xlabel("时刻 (h)")
    ax.set_ylabel("能量 (kWh / 时段)")
    ax.set_xlim(0, 24)
    ax.grid(True, alpha=0.3)
    ax.legend(ncol=3, fontsize=8)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "fig1_dispatch.png", dpi=200)
    plt.close(fig)

    fig, ax1 = plt.subplots(figsize=(10, 4.2))
    ax1.plot(np.concatenate([[0.0], hours]), res.E, color="#1D4ED8", label="SOC (kWh)")
    ax1.set_xlabel("时刻 (h)")
    ax1.set_ylabel("SOC (kWh)")
    ax1.set_xlim(0, 24)
    ax1.set_ylim(1000, 11000)
    ax2 = ax1.twinx()
    ax2.step(hours, data["price"], where="pre", color="#B45309", label="电价 (元/kWh)", alpha=0.85)
    ax2.set_ylabel("电价 (元/kWh)")
    h1, l1 = ax1.get_legend_handles_labels()
    h2, l2 = ax2.get_legend_handles_labels()
    ax1.legend(h1 + h2, l1 + l2, loc="upper right", fontsize=8)
    ax1.grid(True, alpha=0.3)
    fig.tight_layout()
    fig.savefig(FIG_DIR / "fig1_soc_price.png", dpi=200)
    plt.close(fig)
