from __future__ import annotations

from datetime import datetime, time
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from q1.config import (
    CHARGE_HEADERS,
    CHARGE_SHEET,
    DEFAULT_ATTACH1,
    DEFAULT_RESULT1_TEMPLATE,
    DELTA_H,
    FOUR_HOUR_BLOCKS,
    PURCHASE_HEADERS,
    PURCHASE_SHEET,
    T,
)


def normalize_attach_time(value) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M:%S")
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    return str(value).strip()


def load_attachment1(path: Path | None = None) -> pd.DataFrame:
    """Preserve Excel raw row order. Do not sort or rotate."""
    path = Path(path or DEFAULT_ATTACH1)
    if not path.exists():
        raise FileNotFoundError(
            f"附件1 not found: {path}. Set CUMCM_C_ATTACH1 to the xlsx path."
        )
    df = pd.read_excel(path, sheet_name="Sheet1")
    expected = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    if list(df.columns) != expected:
        raise ValueError(f"Unexpected 附件1 columns: {list(df.columns)}")
    if len(df) != T:
        raise ValueError(f"附件1 must have {T} rows, got {len(df)}")

    times = df["时间"].map(normalize_attach_time)
    price = pd.to_numeric(df["电价"], errors="coerce").to_numpy(dtype=float)
    load_kw = pd.to_numeric(df["小区负载"], errors="coerce").to_numpy(dtype=float)
    pv_kw = pd.to_numeric(df["光伏发电预测功率"], errors="coerce").to_numpy(dtype=float)

    if not np.isfinite(price).all() or not np.isfinite(load_kw).all() or not np.isfinite(pv_kw).all():
        raise ValueError("附件1 contains non-finite numeric values")
    if (load_kw < 0).any() or (pv_kw < 0).any():
        raise ValueError("附件1 has negative load or PV")
    if (price < 0).any():
        raise ValueError(
            "附件1 has negative prices; simultaneous charge/discharge logic must be revisited"
        )

    if times.iloc[0] != "00:10:00":
        raise ValueError(f"First attach time must be 00:10:00, got {times.iloc[0]!r}")
    if times.iloc[-1] != "0:00+1":
        raise ValueError(f"Last attach time must be 0:00+1, got {times.iloc[-1]!r}")

    out = pd.DataFrame(
        {
            "t": np.arange(1, T + 1),
            "attach_time": times.to_numpy(),
            "price": price,
            "load_kw": load_kw,
            "pv_kw": pv_kw,
        }
    )
    out["load_kwh"] = out["load_kw"] * DELTA_H
    out["pv_kwh"] = out["pv_kw"] * DELTA_H
    out.attrs["row_order_policy"] = "attachment_raw_order"
    return out


def inspect_template(path: Path | None = None) -> tuple[list[str], list[str]]:
    path = Path(path or DEFAULT_RESULT1_TEMPLATE)
    if not path.exists():
        raise FileNotFoundError(
            f"result1 template not found: {path}. Set CUMCM_C_RESULT1_TEMPLATE."
        )
    wb = load_workbook(path)
    if wb.sheetnames != [PURCHASE_SHEET, CHARGE_SHEET]:
        raise ValueError(f"Unexpected sheets: {wb.sheetnames}")
    ws_g = wb[PURCHASE_SHEET]
    headers_g = (ws_g.cell(1, 1).value, ws_g.cell(1, 2).value)
    if headers_g != PURCHASE_HEADERS:
        raise ValueError(f"计划购电量表头不符: {headers_g}")
    labels = [ws_g.cell(row=i, column=1).value for i in range(2, 2 + T)]
    if any(x is None or str(x).strip() == "" for x in labels):
        raise ValueError("result1 template has empty time labels")
    labels = [str(x) for x in labels]
    if labels[0] != "0:10-0:20" or labels[-1] != "0:00+1-0:10+1":
        raise ValueError(f"Unexpected purchase labels: {labels[0]!r} .. {labels[-1]!r}")

    ws_cd = wb[CHARGE_SHEET]
    headers_cd = tuple(ws_cd.cell(1, c).value for c in range(1, 6))
    if headers_cd != CHARGE_HEADERS:
        raise ValueError(f"充放电量表头不符: {headers_cd}")
    block_labels = [str(ws_cd.cell(row=i, column=1).value) for i in range(2, 8)]
    expected_blocks = [name for name, _a, _b in FOUR_HOUR_BLOCKS]
    if block_labels != expected_blocks:
        raise ValueError(f"四小时区间标签不符: {block_labels}")
    return labels, block_labels


def load_template_labels(path: Path | None = None) -> list[str]:
    labels, _blocks = inspect_template(path)
    return labels
