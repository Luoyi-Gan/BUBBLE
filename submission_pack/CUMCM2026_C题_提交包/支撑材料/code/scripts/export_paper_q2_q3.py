#!/usr/bin/env python3
"""Export Q2/Q3 paper tables and figures from signed-off workbooks and ledgers.

Reads only accepted outputs:
- Q2: `output/result2.xlsx`（V2 政策一致）与 `output/q2_policy_consistent/` 年度台账
- Q3: output/result3.xlsx and output/q3_full_annual/*.csv

Writes LaTeX tables under paper/manuscript/ and PNG/PDF figures under paper/overleaf/.
"""

from __future__ import annotations

import json
import sys
from datetime import datetime, time
from pathlib import Path

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib import font_manager
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

PAPER = ROOT / "paper" / "overleaf"
MANUSCRIPT = ROOT / "paper" / "manuscript"
TABLES_Q2 = MANUSCRIPT / "tables" / "q2"
TABLES_Q3 = MANUSCRIPT / "tables" / "q3"
FIG_Q2 = PAPER / "figures" / "q2"
FIG_Q3 = PAPER / "figures" / "q3"

RESULT2 = ROOT / "output" / "result2.xlsx"
RESULT3 = ROOT / "output" / "result3.xlsx"
Q2_DAILY = ROOT / "output" / "q2_policy_consistent" / "feb_dec_daily_summary.csv"
Q2_COST = ROOT / "output" / "q2_policy_consistent" / "cost_breakdown.csv"
Q2_K = ROOT / "output" / "q2_policy_consistent" / "k_sensitivity.csv"
Q2_CAL = ROOT / "output" / "q2_policy_consistent" / "frozen_policy_calendar.csv"
Q3_STRATEGY = ROOT / "output" / "q3_full_annual" / "q3_annual_strategy_comparison.csv"
Q3_MONTHLY = ROOT / "output" / "q3_full_annual" / "q3_annual_monthly.csv"
Q3_UPDATE = ROOT / "output" / "q3_full_annual" / "q3_update_log.csv"

SAMPLE_DATE = "2025-06-21"
SAMPLE_SLOTS = (
    "10:00-10:10",
    "12:00-12:10",
    "14:00-14:10",
    "16:00-16:10",
    "18:00-18:10",
    "20:00-20:10",
)

BLUE = "#276FBF"
RED = "#C0392B"
GRAY = "#6B7280"
GREEN = "#2F6F4E"


def fmt_num(value: float, digits: int = 4) -> str:
    return f"{float(value):.{digits}f}"


def latex_code(text: str) -> str:
    escaped = text.replace("_", "\\_")
    return f"\\texttt{{{escaped}}}"


def latex_escape(text: str) -> str:
    return (
        text.replace("\\", "\\textbackslash{}")
        .replace("_", "\\_")
        .replace("&", "\\&")
        .replace("%", "\\%")
    )


def configure_style() -> str:
    available = {font.name for font in font_manager.fontManager.ttflist}
    candidates = (
        "Noto Sans CJK SC",
        "PingFang SC",
        "Hiragino Sans GB",
        "Heiti SC",
        "Songti SC",
        "Arial Unicode MS",
    )
    font = next((name for name in candidates if name in available), "DejaVu Sans")
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [font],
            "axes.unicode_minus": False,
            "figure.facecolor": "white",
            "axes.facecolor": "white",
            "font.size": 10,
            "axes.titlesize": 12,
            "axes.labelsize": 10,
            "legend.fontsize": 9,
            "pdf.fonttype": 42,
        }
    )
    return font


def save_figure(fig: plt.Figure, stem: str, out_dir: Path) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_dir / f"{stem}.pdf", bbox_inches="tight")
    fig.savefig(out_dir / f"{stem}.png", dpi=320, bbox_inches="tight")
    plt.close(fig)


def write_text(path: Path, content: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(content, encoding="utf-8")


def load_purchase_row(path: Path, sheet: str, date_text: str) -> tuple[list[str], dict[str, float], float, float]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows)]
    target = pd.Timestamp(date_text).to_pydatetime()
    values: dict[str, float] = {}
    day_total = 0.0
    day_fee = 0.0
    for row in rows:
        if row[0] is None:
            continue
        if pd.Timestamp(row[0]).normalize() != pd.Timestamp(target).normalize():
            continue
        for idx, header in enumerate(headers[1:], start=1):
            if idx >= len(row):
                break
            if header in SAMPLE_SLOTS:
                values[header] = float(row[idx] or 0.0)
        if len(row) > 145:
            day_total = float(row[145] or 0.0)
        if len(row) > 146:
            day_fee = float(row[146] or 0.0)
        break
    wb.close()
    if not values:
        raise ValueError(f"missing purchase row for {date_text} in {path}")
    return headers, values, day_total, day_fee


def load_charge_blocks(path: Path, date_text: str) -> tuple[list[tuple[str, float, float]], float, float]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["充放电量"]
    blocks: list[tuple[str, float, float]] = []
    soc_start = soc_end = 0.0
    target = pd.Timestamp(date_text).normalize()
    for row in ws.iter_rows(values_only=True):
        if row[0] == "日期" or row[1] == "时间段":
            continue
        if row[0] is not None:
            try:
                row_date = pd.Timestamp(row[0]).normalize()
            except (ValueError, TypeError):
                continue
            if row_date == target:
                soc_start = float(row[5] or 0.0)
        if row[1] is None:
            continue
        if row[0] is not None:
            try:
                row_date = pd.Timestamp(row[0]).normalize()
            except (ValueError, TypeError):
                row_date = None
            if row_date == target:
                blocks.append((str(row[1]), float(row[2] or 0.0), float(row[3] or 0.0)))
                if str(row[4]) in ("24:00", "0:00+1") or row[4] == time(0, 0):
                    soc_end = float(row[5] or 0.0)
        elif row[0] is None and blocks and len(blocks) < 6:
            blocks.append((str(row[1]), float(row[2] or 0.0), float(row[3] or 0.0)))
            if str(row[4]) in ("24:00", "0:00+1"):
                soc_end = float(row[5] or 0.0)
    wb.close()
    if len(blocks) != 6:
        raise ValueError(f"expected 6 charge blocks for {date_text}, got {len(blocks)}")
    return blocks, soc_start, soc_end


def load_emergency_segments(path: Path, date_text: str) -> list[tuple[str, float]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["紧急购电量"]
    target = pd.Timestamp(date_text).normalize()
    segments: list[tuple[str, float]] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        if pd.Timestamp(row[0]).normalize() != target:
            continue
        segments.append((str(row[1]), float(row[2] or 0.0)))
    wb.close()
    return segments


def load_adjust_row(path: Path, date_text: str) -> dict[str, float]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["调整购电量"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows)]
    target = pd.Timestamp(date_text).normalize()
    values: dict[str, float] = {}
    for row in rows:
        if row[0] is None:
            continue
        if pd.Timestamp(row[0]).normalize() != target:
            continue
        for idx, header in enumerate(headers[1:], start=1):
            if header in SAMPLE_SLOTS:
                values[header] = float(row[idx] or 0.0)
        break
    wb.close()
    if not values:
        raise ValueError(f"missing adjust row for {date_text}")
    return values


def render_triplet_table(
    caption: str,
    label: str,
    value_title: str,
    rows: list[tuple[str, str]],
    footer: tuple[str, str, str, str] | None = None,
) -> str:
    lines = [
        "\\begin{table}[htbp]",
        "\\centering",
        f"\\caption{{{caption}}}\\label{{{label}}}",
        "\\normalsize\\setlength{\\tabcolsep}{5pt}",
        "\\begin{tabular}{@{}cccccc@{}}",
        "\\toprule",
        f"时间段 & {value_title} & 时间段 & {value_title} & 时间段 & {value_title} \\\\",
        "\\midrule",
    ]
    for i in range(0, len(rows), 3):
        chunk = rows[i : i + 3]
        while len(chunk) < 3:
            chunk.append(("", ""))
        lines.append(
            f"{chunk[0][0]} & {chunk[0][1]} & {chunk[1][0]} & {chunk[1][1]} & "
            f"{chunk[2][0]} & {chunk[2][1]} \\\\"
        )
    if footer:
        left_label, left_val, right_label, right_val = footer
        lines.append(
            f"{left_label} & \\multicolumn{{2}}{{c}}{{{left_val}}} & "
            f"{right_label} & \\multicolumn{{2}}{{c}}{{{right_val}}} \\\\"
        )
    lines.extend(["\\bottomrule", "\\end{tabular}", "\\end{table}", ""])
    return "\n".join(lines)


def export_q2_tables() -> dict:
    _, purchase_vals, day_q, day_fee = load_purchase_row(RESULT2, "计划购电量", SAMPLE_DATE)
    blocks, soc_start, soc_end = load_charge_blocks(RESULT2, SAMPLE_DATE)
    emergency = load_emergency_segments(RESULT2, SAMPLE_DATE)
    cost = pd.read_csv(Q2_COST)
    k_df = pd.read_csv(Q2_K)
    feb = cost[cost["label"] == "feb_dec_inherited_warmup_soc"].iloc[0]
    jan = cost[cost["label"] == "january_warmup_main"].iloc[0]

    purchase_rows = [(slot, fmt_num(purchase_vals[slot])) for slot in SAMPLE_SLOTS]
    purchase_tex = render_triplet_table(
        f"问题二 {SAMPLE_DATE} 计划购电量及全天汇总（政策一致 $K=8$ 主方案）",
        "tab:q2-purchase",
        "购电量",
        purchase_rows,
        footer=("全天购电量", fmt_num(day_q), "全天购电费", fmt_num(day_fee)),
    )

    storage_lines = [
        "\\begin{table}[htbp]",
        "\\centering",
        f"\\caption{{问题二 {SAMPLE_DATE} 充放电量及日初/日末储电量}}\\label{{tab:q2-storage}}",
        "\\normalsize\\setlength{\\tabcolsep}{5pt}",
        "\\begin{tabular}{@{}cccccc@{}}",
        "\\toprule",
        "时间段 & 充电量 & 放电量 & 时间段 & 充电量 & 放电量 \\\\",
        "\\midrule",
    ]
    for i in range(0, 6, 2):
        b0, b1 = blocks[i], blocks[i + 1]
        storage_lines.append(
            f"{b0[0]} & {fmt_num(b0[1])} & {fmt_num(b0[2])} & "
            f"{b1[0]} & {fmt_num(b1[1])} & {fmt_num(b1[2])} \\\\"
        )
    storage_lines.extend(
        [
            f"0:00 储电量 & \\multicolumn{{2}}{{c}}{{{fmt_num(soc_start)}}} & "
            f"24:00 储电量 & \\multicolumn{{2}}{{c}}{{{fmt_num(soc_end)}}} \\\\",
            "\\bottomrule",
            "\\end{tabular}",
            "\\end{table}",
            "",
        ]
    )
    storage_tex = "\n".join(storage_lines)

    emergency_lines = [
        "\\begin{table}[htbp]",
        "\\centering",
        f"\\caption{{问题二 {SAMPLE_DATE} 紧急购电记录（2--12 月输出区间内典型日）}}\\label{{tab:q2-emergency}}",
        "\\normalsize\\setlength{\\tabcolsep}{6pt}",
        "\\begin{tabular}{@{}ccc@{}}",
        "\\toprule",
        "购电时间段 & 购电量 & 备注 \\\\",
        "\\midrule",
    ]
    total_e = 0.0
    for interval, amount in emergency[:8]:
        total_e += amount
        emergency_lines.append(f"{interval} & {fmt_num(amount)} & 紧急购电 \\\\")
    if len(emergency) > 8:
        emergency_lines.append(
            f"\\multicolumn{{3}}{{c}}{{其余 {len(emergency) - 8} 条记录见 \\texttt{{result2.xlsx}}}} \\\\"
        )
    emergency_lines.extend(
        [
            f"\\midrule",
            f"当日合计 & {fmt_num(sum(x for _, x in emergency))} & kWh \\\\",
            "\\bottomrule",
            "\\end{tabular}",
            "\\end{table}",
            "",
        ]
    )
    emergency_tex = "\n".join(emergency_lines)

    cost_tex = f"""\\begin{{table}}[htbp]
\\centering
\\caption{{问题二年度成本分解（政策一致 $K=8$；2--12 月为题设输出区间）}}\\label{{tab:q2-cost-summary}}
\\normalsize\\setlength{{\\tabcolsep}}{{6pt}}
\\begin{{tabular}}{{@{{}}lrrrr@{{}}}}
\\toprule
统计口径 & 天数 & 计划购电费/元 & 紧急购电费/元 & 总成本/元 \\\\
\\midrule
1 月预热（不计入题设输出） & 31 & {fmt_num(jan.planned_cost_yuan, 2)} & {fmt_num(jan.emergency_cost_yuan, 2)} & {fmt_num(jan.total_cost_yuan, 2)} \\\\
2--12 月正式输出 & 334 & {fmt_num(feb.planned_cost_yuan, 2)} & {fmt_num(feb.emergency_cost_yuan, 2)} & {fmt_num(feb.total_cost_yuan, 2)} \\\\
\\midrule
2--12 月紧急购电量/kWh & \\multicolumn{{4}}{{c}}{{{fmt_num(feb.emergency_kwh, 2)}}} \\\\
2 月 1 日继承 SOC/kWh & \\multicolumn{{4}}{{c}}{{{fmt_num(feb.soc_start_kwh, 4)}}} \\\\
\\bottomrule
\\end{{tabular}}
\\end{{table}}
"""

    k_rows = []
    for _, row in k_df.iterrows():
        k_rows.append(
            f"$K={int(row.scenario_k)}$ & {int(row.n_days)} & "
            f"{fmt_num(row.total_cost_yuan, 2)} & {fmt_num(row.emergency_kwh, 2)} & "
            f"{fmt_num(row.elapsed_seconds, 1)} & {latex_code(str(row.conclusion))} \\\\"
        )
    param_tex = f"""\\begin{{table}}[htbp]
\\centering
\\caption{{问题二情景数 $K$ 的闭环复核（2--12 月输出区间；正式主方案取 $K=8$）}}\\label{{tab:q2-parameter}}
\\normalsize\\setlength{{\\tabcolsep}}{{5pt}}
\\begin{{tabular}}{{@{{}}crrrrl@{{}}}}
\\toprule
$K$ & 天数 & 总成本/元 & 紧急购电/kWh & 耗时/s & 结论 \\\\
\\midrule
{chr(10).join(k_rows)}\\bottomrule
\\end{{tabular}}
\\end{{table}}
"""

    write_text(TABLES_Q2 / "purchase.tex", purchase_tex)
    write_text(TABLES_Q2 / "storage.tex", storage_tex)
    write_text(TABLES_Q2 / "emergency.tex", emergency_tex)
    write_text(TABLES_Q2 / "cost_summary.tex", cost_tex)
    write_text(TABLES_Q2 / "parameter_selection.tex", param_tex)
    return {
        "sample_date": SAMPLE_DATE,
        "feb_dec_total_yuan": float(feb.total_cost_yuan),
        "emergency_segments": len(emergency),
    }


def export_q3_tables() -> dict:
    _, g0_vals, day_g0, _ = load_purchase_row(RESULT3, "计划购电量", SAMPLE_DATE)
    g_vals = load_adjust_row(RESULT3, SAMPLE_DATE)
    blocks, soc_start, soc_end = load_charge_blocks(RESULT3, SAMPLE_DATE)
    emergency = load_emergency_segments(RESULT3, SAMPLE_DATE)
    strategy = pd.read_csv(Q3_STRATEGY)
    a = strategy[
        (strategy["strategy"].isin(["M0", "M1_M6"]))
        & (strategy["year_end_boundary"] == "A_q2_aligned")
    ].set_index("strategy")

    def build_value_table(caption: str, label: str, col_title: str, mapping: dict[str, float]) -> str:
        rows = [(slot, fmt_num(mapping[slot])) for slot in SAMPLE_SLOTS]
        lines = [
            "\\begin{table}[htbp]",
            "\\centering",
            f"\\caption{{{caption}}}\\label{{{label}}}",
            "\\normalsize\\setlength{\\tabcolsep}{5pt}",
            "\\begin{tabular}{@{}cccccc@{}}",
            "\\toprule",
            f"时间段 & {col_title} & 时间段 & {col_title} & 时间段 & {col_title} \\\\",
            "\\midrule",
        ]
        for i in range(0, 6, 3):
            chunk = rows[i : i + 3]
            lines.append(
                f"{chunk[0][0]} & {chunk[0][1]} & {chunk[1][0]} & {chunk[1][1]} & "
                f"{chunk[2][0]} & {chunk[2][1]} \\\\"
            )
        lines.extend(["\\bottomrule", "\\end{tabular}", "\\end{table}", ""])
        return "\n".join(lines)

    purchase_tex = build_value_table(
        f"问题三 {SAMPLE_DATE} 计划购电量 $g^0$（A 边界 + M1\\_M6 主方案）",
        "tab:q3-purchase",
        "计划购电量",
        g0_vals,
    )
    adjust_tex = build_value_table(
        f"问题三 {SAMPLE_DATE} 调整购电量 $g^F$（A 边界 + M1\\_M6 主方案）",
        "tab:q3-adjust",
        "调整购电量",
        g_vals,
    )

    storage_lines = [
        "\\begin{table}[htbp]",
        "\\centering",
        f"\\caption{{问题三 {SAMPLE_DATE} 充放电量及日初/日末储电量}}\\label{{tab:q3-storage}}",
        "\\normalsize\\setlength{\\tabcolsep}{5pt}",
        "\\begin{tabular}{@{}cccccc@{}}",
        "\\toprule",
        "时间段 & 充电量 & 放电量 & 时间段 & 充电量 & 放电量 \\\\",
        "\\midrule",
    ]
    for i in range(0, 6, 2):
        b0, b1 = blocks[i], blocks[i + 1]
        storage_lines.append(
            f"{b0[0]} & {fmt_num(b0[1])} & {fmt_num(b0[2])} & "
            f"{b1[0]} & {fmt_num(b1[1])} & {fmt_num(b1[2])} \\\\"
        )
    storage_lines.extend(
        [
            f"0:00 储电量 & \\multicolumn{{2}}{{c}}{{{fmt_num(soc_start)}}} & "
            f"24:00 储电量 & \\multicolumn{{2}}{{c}}{{{fmt_num(soc_end)}}} \\\\",
            "\\bottomrule",
            "\\end{tabular}",
            "\\end{table}",
            "",
        ]
    )

    emergency_lines = [
        "\\begin{table}[htbp]",
        "\\centering",
        f"\\caption{{问题三 {SAMPLE_DATE} 紧急购电记录}}\\label{{tab:q3-emergency}}",
        "\\normalsize\\setlength{\\tabcolsep}{6pt}",
        "\\begin{tabular}{@{}ccc@{}}",
        "\\toprule",
        "购电时间段 & 购电量 & 备注 \\\\",
        "\\midrule",
    ]
    for interval, amount in emergency[:6]:
        emergency_lines.append(f"{interval} & {fmt_num(amount)} & 紧急购电 \\\\")
    if len(emergency) > 6:
        emergency_lines.append(
            f"\\multicolumn{{3}}{{c}}{{其余 {len(emergency) - 6} 条记录见 \\texttt{{result3.xlsx}}}} \\\\"
        )
    emergency_lines.extend(
        [
            "\\midrule",
            f"当日合计 & {fmt_num(sum(x for _, x in emergency))} & kWh \\\\",
            "\\bottomrule",
            "\\end{tabular}",
            "\\end{table}",
            "",
        ]
    )

    m0 = a.loc["M0"]
    m16 = a.loc["M1_M6"]
    saving = float(m0.total_cost_yuan - m16.total_cost_yuan)
    saving_pct = 100.0 * saving / float(m0.total_cost_yuan)
    compare_tex = f"""\\begin{{table}}[htbp]
\\centering
\\caption{{问题三全年策略比较（A 边界；1--12 月连续运行口径）}}\\label{{tab:q3-strategy-compare}}
\\normalsize\\setlength{{\\tabcolsep}}{{5pt}}
\\begin{{tabular}}{{@{{}}lrrrrr@{{}}}}
\\toprule
策略 & 总成本/元 & 结算成本/元 & 紧急购电/元 & 紧急购电/kWh & 调整次数 \\\\
\\midrule
M0（不调整） & {fmt_num(m0.total_cost_yuan, 2)} & {fmt_num(m0.settlement_cost_yuan, 2)} & {fmt_num(m0.emergency_cost_yuan, 2)} & {fmt_num(m0.emergency_kwh, 2)} & {int(m0.adjustment_count)} \\\\
M1\\_M6（主方案） & {fmt_num(m16.total_cost_yuan, 2)} & {fmt_num(m16.settlement_cost_yuan, 2)} & {fmt_num(m16.emergency_cost_yuan, 2)} & {fmt_num(m16.emergency_kwh, 2)} & {int(m16.adjustment_count)} \\\\
\\midrule
相对 M0 节省 & \\multicolumn{{5}}{{c}}{{{fmt_num(saving, 2)} 元（{fmt_num(saving_pct, 2)}\\%）}} \\\\
正式 result3 输出区间成本 & \\multicolumn{{5}}{{c}}{{14512748.53 元（2025-02-01 至 2025-12-31，非全年口径）}} \\\\
\\bottomrule
\\end{{tabular}}
\\end{{table}}
"""

    write_text(TABLES_Q3 / "purchase.tex", purchase_tex)
    write_text(TABLES_Q3 / "adjust.tex", adjust_tex)
    write_text(TABLES_Q3 / "storage.tex", "\n".join(storage_lines))
    write_text(TABLES_Q3 / "emergency.tex", "\n".join(emergency_lines))
    write_text(TABLES_Q3 / "strategy_comparison.tex", compare_tex)
    return {
        "annual_m1m6_yuan": float(m16.total_cost_yuan),
        "annual_saving_yuan": saving,
        "export_window_yuan": 14512748.527966088,
    }


def export_q2_figures() -> None:
    configure_style()
    daily = pd.read_csv(Q2_DAILY, parse_dates=["date"])
    daily = daily[(daily["date"] >= "2025-02-01") & (daily["date"] <= "2025-12-31")]
    daily["month"] = daily["date"].dt.to_period("M").astype(str)
    monthly = daily.groupby("month", as_index=False).agg(
        planned_cost_yuan=("planned_cost_yuan", "sum"),
        emergency_cost_yuan=("emergency_cost_yuan", "sum"),
    )
    x = np.arange(len(monthly))
    fig, ax = plt.subplots(figsize=(11.5, 5.2))
    ax.bar(x, monthly["planned_cost_yuan"] / 1e6, label="计划购电费", color=BLUE)
    ax.bar(
        x,
        monthly["emergency_cost_yuan"] / 1e6,
        bottom=monthly["planned_cost_yuan"] / 1e6,
        label="紧急购电费",
        color=RED,
        alpha=0.85,
    )
    ax.set_xticks(x, [m.replace("2025-", "") + "月" for m in monthly["month"]], rotation=0)
    ax.set_ylabel("成本 / 百万元")
    ax.set_title("问题二 2025 年 2--12 月月度购电成本构成（政策一致 $K=8$）")
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    save_figure(fig, "fig_q2_monthly_cost", FIG_Q2)

    fig, ax = plt.subplots(figsize=(11.5, 4.8))
    ax.plot(daily["date"], daily["soc_end_kwh"], color=BLUE, linewidth=0.9)
    ax.axhline(1200, color=GRAY, linestyle="--", linewidth=0.8, label="SOC 下界")
    ax.axhline(10800, color=GRAY, linestyle=":", linewidth=0.8, label="SOC 上界")
    ax.set_ylabel("日末 SOC / kWh")
    ax.set_title("问题二 2025 年 2--12 月日末储电量轨迹")
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m月"))
    ax.legend(frameon=False, loc="upper right")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    save_figure(fig, "fig_q2_soc_daily", FIG_Q2)

    cal = pd.read_csv(Q2_CAL, parse_dates=["date"])
    selected = (
        cal.groupby(["effective_start_date", "effective_end_date", "forecast_mode", "risk_alpha"], dropna=False)
        .size()
        .reset_index(name="n_days")
        .sort_values("effective_start_date")
    )
    fig, ax = plt.subplots(figsize=(11.5, 3.8))
    for _, row in selected.iterrows():
        start = pd.Timestamp(row.effective_start_date)
        end = pd.Timestamp(row.effective_end_date) + pd.Timedelta(days=1)
        alpha = "" if pd.isna(row.risk_alpha) else f", α={row.risk_alpha:.2f}"
        label = f"{row.forecast_mode}{alpha}"
        ax.barh(label, (end - start).days, left=start, height=0.55, color=BLUE, alpha=0.75)
    ax.axvspan(pd.Timestamp("2025-01-01"), pd.Timestamp("2025-02-01"), color=GRAY, alpha=0.12)
    ax.set_xlim(pd.Timestamp("2025-01-01"), pd.Timestamp("2026-01-01"))
    ax.xaxis.set_major_locator(mdates.MonthLocator())
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%m月"))
    ax.set_title("问题二 C2-R2 冻结预测结构与风险分位日历（$K=8$ 主方案沿用）")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    save_figure(fig, "fig_q2_policy_calendar", FIG_Q2)


def export_q3_figures() -> None:
    configure_style()
    strategy = pd.read_csv(Q3_STRATEGY)
    subset = strategy[
        (strategy["year_end_boundary"] == "A_q2_aligned")
        & (strategy["strategy"].isin(["M0", "M1_M6", "M6_only", "M12_only"]))
    ].copy()
    subset = subset.sort_values("total_cost_yuan")
    labels = subset["strategy"].tolist()
    planned = subset["settlement_cost_yuan"].to_numpy()
    emergency = subset["emergency_cost_yuan"].to_numpy()
    x = np.arange(len(labels))
    fig, ax = plt.subplots(figsize=(8.8, 5.0))
    ax.bar(x, planned / 1e6, label="结算成本", color=BLUE)
    ax.bar(x, emergency / 1e6, bottom=planned / 1e6, label="紧急购电成本", color=RED, alpha=0.85)
    ax.set_xticks(x, labels, rotation=15)
    ax.set_ylabel("成本 / 百万元")
    ax.set_title("问题三 A 边界下各策略全年总成本分解")
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    save_figure(fig, "fig_q3_strategy_cost", FIG_Q3)

    monthly = pd.read_csv(Q3_MONTHLY)
    monthly = monthly[monthly["year_end_boundary"] == "A_q2_aligned"]
    months = sorted(m for m in monthly["month"].unique() if m >= "2025-02")
    m0 = monthly[(monthly["strategy"] == "M0")].set_index("month").loc[months]["emergency_cost_yuan"]
    m16 = monthly[(monthly["strategy"] == "M1_M6")].set_index("month").loc[months]["emergency_cost_yuan"]
    x = np.arange(len(months))
    width = 0.36
    fig, ax = plt.subplots(figsize=(11.5, 5.0))
    ax.bar(x - width / 2, m0 / 1e4, width=width, label="M0", color=GRAY)
    ax.bar(x + width / 2, m16 / 1e4, width=width, label="M1_M6", color=BLUE)
    ax.set_xticks(x, [m.replace("2025-", "") + "月" for m in months])
    ax.set_ylabel("紧急购电成本 / 万元")
    ax.set_title("问题三 2025 年 2--12 月紧急购电成本月度对比")
    ax.legend(frameon=False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    save_figure(fig, "fig_q3_monthly_emergency", FIG_Q3)

    updates = pd.read_csv(Q3_UPDATE, parse_dates=["date"])
    updates = updates[
        (updates["strategy"] == "M1_M6")
        & (updates["year_end_boundary"] == "A_q2_aligned")
        & (updates["update_time"].isin(["06:00", "12:00", "18:00"]))
        & (updates["implemented"] == True)
    ].copy()
    updates["month"] = updates["date"].dt.to_period("M").astype(str)
    monthly_voi = updates.groupby("month")["voi_yuan"].sum()
    months_voi = [m for m in sorted(monthly_voi.index) if m >= "2025-02"]
    fig, ax = plt.subplots(figsize=(11.5, 4.6))
    ax.bar(
        np.arange(len(months_voi)),
        [monthly_voi[m] / 1e4 for m in months_voi],
        color=GREEN,
        alpha=0.85,
    )
    ax.set_xticks(np.arange(len(months_voi)), [m.replace("2025-", "") + "月" for m in months_voi])
    ax.set_ylabel("信息价值 / 万元")
    ax.set_title("问题三 M1\\_M6 月度已实施调整的信息经济价值（VoI 汇总）")
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    save_figure(fig, "fig_q3_update_value", FIG_Q3)


def main() -> None:
    for path in (RESULT2, RESULT3, Q2_DAILY, Q2_COST, Q3_STRATEGY):
        if not path.exists():
            raise FileNotFoundError(path)
    q2_meta = export_q2_tables()
    q3_meta = export_q3_tables()
    export_q2_figures()
    export_q3_figures()
    audit = {"q2": q2_meta, "q3": q3_meta}
    write_text(PAPER / "paper_assets_audit.json", json.dumps(audit, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps(audit, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
