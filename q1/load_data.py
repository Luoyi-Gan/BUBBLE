from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from q1.config import DEFAULT_ATTACH1, DEFAULT_RESULT1_TEMPLATE, DELTA_H, T


def load_attachment1(path: Path | None = None) -> pd.DataFrame:
    path = Path(path or DEFAULT_ATTACH1)
    if not path.exists():
        raise FileNotFoundError(
            f"附件1 not found: {path}. Set CUMCM_C_ATTACH1 to the xlsx path."
        )
    df = pd.read_excel(path, sheet_name="Sheet1")
    expected = ["时间", "电价", "小区负载", "光伏发电预测功率"]
    if list(df.columns) != expected:
        raise ValueError(f"Unexpected 附件1 columns: {list(df.columns)}")
    if len(df) != T:
        raise ValueError(f"附件1 must have {T} rows, got {len(df)}")
    if df.isna().any().any():
        raise ValueError("附件1 contains missing values")
    if (df["小区负载"] < 0).any() or (df["光伏发电预测功率"] < 0).any():
        raise ValueError("附件1 has negative load or PV")

    out = pd.DataFrame(
        {
            "t": np.arange(1, T + 1),
            "attach_time": df["时间"].astype(str),
            "price": df["电价"].to_numpy(dtype=float),
            "load_kw": df["小区负载"].to_numpy(dtype=float),
            "pv_kw": df["光伏发电预测功率"].to_numpy(dtype=float),
        }
    )
    out["load_kwh"] = out["load_kw"] * DELTA_H
    out["pv_kwh"] = out["pv_kw"] * DELTA_H
    return out


def load_template_labels(path: Path | None = None) -> list[str]:
    path = Path(path or DEFAULT_RESULT1_TEMPLATE)
    if not path.exists():
        raise FileNotFoundError(
            f"result1 template not found: {path}. Set CUMCM_C_RESULT1_TEMPLATE."
        )
    wb = load_workbook(path, data_only=True)
    ws = wb["计划购电量"]
    labels = [ws.cell(row=i, column=1).value for i in range(2, 2 + T)]
    if any(x is None or str(x).strip() == "" for x in labels):
        raise ValueError("result1 template has empty time labels")
    if len(labels) != T:
        raise ValueError(f"result1 template must have {T} time labels")
    return [str(x) for x in labels]
